"""
Preço site por EAN  —  v9

Novo na v9:
  • vários sites ao mesmo tempo (--paralelo, padrão 6). Cada site continua no seu ritmo e com no máximo
    uma consulta em andamento — para o site nada muda; a rodada termina em fração do tempo.
  • Excel parcial (preco_site_PARCIAL.xlsx) regravado a cada 15 min durante a rodada.
  • fuso de Brasília automático (o Colab roda em UTC e datava a coleta noturna como "amanhã").

Aceita dois tipos de planilha (detecta sozinho):
  • RÉGUA DE PREÇOS (ex.: "Busca preços - Atualizado.xlsx", aba NOVA RÉGUA): lista de EANs SEM cliente.
    Cada EAN é consultado em TODOS os clientes configurados em SITES, e o Excel compara o preço do
    site com o Preço Sugerido (coluna P da régua). Linhas sem EAN são descartadas.
  • ACELERA (CLIENTE | EAN | SKU): consulta cada EAN só no cliente da linha (modo antigo).

Instalação (uma vez):
  pip install pandas openpyxl requests beautifulsoup4 playwright
  python -m playwright install chromium
  # Linux / Colab, se reclamar de biblioteca do sistema:
  python -m playwright install-deps chromium

No Colab: rode as duas linhas acima ANTES do script (a máquina é nova a cada sessão e vem
sem navegador). Sem isso, só os clientes com API (VTEX/SFCC) devolvem preço — o script avisa
no começo e não gasta tempo tentando. '--ver' não funciona no Colab (não há tela).

Dois modos de rodar:
  CONTINUAR (padrão)  consulta só o que falta ou já venceu no cache — pode parar e retomar à vontade
  RODAR TUDO (--tudo) refaz todos os EANs da seleção, ignorando o cache

Uso:
  python preco_site_por_ean.py                        # acha sozinho a régua (Busca*/Régua*) ou o Acelera*.xlsx
  python preco_site_por_ean.py "Busca preços - Atualizado.xlsx"   # régua: EANs x todos os clientes
  python preco_site_por_ean.py acelera.xlsx           # Acelera: EAN x cliente da linha
  python preco_site_por_ean.py --tudo                 # refaz tudo (ignora o cache)
  python preco_site_por_ean.py --so ARAUJO NISSEI     # só esses clientes
  python preco_site_por_ean.py --ver                  # navegador visível (ajuda quando o site bloqueia)
  python preco_site_por_ean.py --mobile               # emula celular (Android/Chrome): site mais simples, menos bloqueio
  python preco_site_por_ean.py --passadas 3           # mais uma rodada em cima do que ficou sem preço
  python preco_site_por_ean.py --limite 5             # só 5 EANs por cliente (teste rápido)
  python preco_site_por_ean.py --paralelo 8           # 8 sites ao mesmo tempo (padrão 6; 1 = um de cada vez)
  python preco_site_por_ean.py --relatorio            # só regera o Excel a partir do cache, sem consultar
  python preco_site_por_ean.py --sem-navegador        # só motores de API (VTEX / SFCC), sem abrir navegador
  python preco_site_por_ean.py --debug                # salva print + HTML de cada página em preco_site_dados/debug/
  python preco_site_por_ean.py --rediagnosticar       # refaz o teste de qual método funciona em cada site
  python preco_site_por_ean.py --listar               # mostra os clientes, bandeiras e métodos
  python preco_site_por_ean.py --testar PAGUE MENOS 7896004707037 [--nome "DIPIRONA 500MG C/30"]
  python preco_site_por_ean.py --mapear-api NISSEI 7896004707037   # mostra de qual JSON o site tira o preço

Como o robô acha o preço de cada EAN:
  1. se já conhece a página daquele produto, vai direto nela (pula a busca, que é a parte bloqueada)
  2. API do cliente (VTEX / SFCC) consultada pelo EAN
  3. JSON que a própria página busca, interceptado pelo navegador — é o que o app do site usa
  4. dados estruturados da página (JSON-LD / microdata com GTIN)
  5. preço no DOM / no texto, com o EAN confirmado na página
  6. não achou pelo EAN? busca pelo nome do SKU; ainda não? procura a página no sitemap do site
  Sem resultado na bandeira principal, tenta as outras bandeiras do mesmo grupo (ex.: Drogasil -> Droga Raia).
  No fim da rodada, quem ficou sem preço volta para uma 2ª passada, mais devagar e por outros caminhos.

Arquivos gerados:
  preco_site_AAAAMMDD_HHMM.xlsx            abas: Preços | Conferir | Resumo | Ignorados
                                           (com régua: + aba "Régua x Clientes" — um EAN por linha, um
                                            cliente por coluna — e Preços/Resumo comparados ao Preço Sugerido)
  preco_site_dados/cache_precos.json       cache com validade por status (evita reconsultar o mesmo EAN)
  preco_site_dados/metodos_clientes.json   método que funcionou em cada site + URLs de busca aprendidas
  preco_site_dados/indice_urls.json        página de cada produto por EAN -> rodadas seguintes vão direto
  preco_site_dados/historico_precos.csv    histórico de preços (uma linha por coleta) -> "preço anterior" no Excel
  preco_site_dados/sitemaps/               mapa de URLs de cada site (renovado a cada 7 dias)
  preco_site_dados/logs/                   log detalhado de cada execução
  preco_site_dados/sites_extra.json        (opcional) clientes/sites extras sem mexer no código:
                                           {"NOVO CLIENTE": {"motor": "auto", "base": "https://www.site.com.br",
                                                             "bandeiras": [{"nome": "Outra", "base": "https://..."}]}}
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import math
import os
import queue
import random
import re
import sys
import threading
import time
import unicodedata
import warnings
from collections import OrderedDict, deque
from dataclasses import asdict, dataclass, fields, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional
from urllib.parse import quote, quote_plus, urljoin, urlparse

try:
    import pandas as pd
    import requests
    from bs4 import BeautifulSoup
except ImportError as e:  # pragma: no cover
    sys.exit(f"\n>>> Falta biblioteca: {e.name}. Rode:\n"
             "    pip install pandas openpyxl requests beautifulsoup4 playwright\n")

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# o Colab (e a maioria dos servidores) roda em UTC: sem isto, depois das 21h a coleta sai com a data de amanhã
if not os.environ.get("TZ") and hasattr(time, "tzset"):
    os.environ["TZ"] = "America/Sao_Paulo"
    time.tzset()

warnings.filterwarnings("ignore", message="Unverified HTTPS request")
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
warnings.filterwarnings("ignore", category=FutureWarning)

# ---------------------------------------------------------------------------
# CLIENTES  (motor: vtex | sfcc | navegador | auto)
#   base   -> home do e-commerce da bandeira principal
#   nome   -> (opcional) como a bandeira aparece no relatório
#   busca  -> (opcional) URL de busca com {termo}; sem ela o robô usa a caixa de busca do site
#             e APRENDE a URL sozinho na primeira busca que der certo
#   ritmo  -> (opcional) segundos mínimos entre consultas nesse site
#   bandeiras -> (opcional) outras lojas do mesmo grupo, na ordem em que devem ser tentadas
#                quando o EAN não aparece na principal (ou quando ela bloqueia)
#
# Conferido em set/2026: cada domínio abaixo foi aberto e confirmado como loja ativa.
# ---------------------------------------------------------------------------
SITES: dict = {
    "PAGUE MENOS":   {"motor": "vtex", "base": "https://www.paguemenos.com.br",
                      "bandeiras": [{"nome": "Extrafarma", "motor": "vtex",
                                     "base": "https://www.extrafarma.com.br"}]},
    "DPSP":          {"motor": "vtex", "base": "https://www.drogariasaopaulo.com.br", "nome": "Drogaria São Paulo",
                      "bandeiras": [{"nome": "Drogarias Pacheco", "motor": "vtex",
                                     "base": "https://www.drogariaspacheco.com.br"}]},
    "SÃO JOÃO":      {"motor": "vtex", "base": "https://www.saojoaofarmacias.com.br"},
    "VENANCIO":      {"motor": "vtex", "base": "https://www.drogariavenancio.com.br"},
    "DROGAL":        {"motor": "vtex", "base": "https://www.drogal.com.br"},
    # drogariaindiana.com.br redireciona (302) para farmaciaindiana.com.br
    "INDIANA":       {"motor": "vtex", "base": "https://www.farmaciaindiana.com.br"},
    "CLAMED":        {"motor": "vtex", "base": "https://www.drogariacatarinense.com.br", "nome": "Drogaria Catarinense",
                      "bandeiras": [{"nome": "Preço Popular", "motor": "auto",
                                     "base": "https://www.precopopular.com.br"}]},

    "ARAUJO":        {"motor": "sfcc", "base": "https://www.araujo.com.br",
                      "busca": "https://www.araujo.com.br/busca?q={termo}", "ritmo": 15},
    # a NISSEI passa o termo pelo caminho da URL, não por query string (a caixa de busca não tem 'name')
    "NISSEI":        {"motor": "navegador", "base": "https://www.farmaciasnissei.com.br",
                      "busca": "https://www.farmaciasnissei.com.br/pesquisa/{termo}"},
    "RAIA DROGASIL": {"motor": "navegador", "base": "https://www.drogasil.com.br", "nome": "Drogasil",
                      "busca": "https://www.drogasil.com.br/search?w={termo}", "ritmo": 20,
                      "bandeiras": [{"nome": "Droga Raia", "motor": "navegador",
                                     "base": "https://www.drogaraia.com.br",
                                     "busca": "https://www.drogaraia.com.br/search?w={termo}"}]},
    "PANVEL":        {"motor": "navegador", "base": "https://www.panvel.com",
                      "busca": "https://www.panvel.com/panvel/buscarProduto.do?termoPesquisa={termo}"},

    # Guedes e Paixão Ltda. é a razão social da Drogaria Minas-Brasil (Montes Claros/MG)
    "GUEDES E PAIXAO":           {"motor": "auto", "base": "https://www.drogariaminasbrasil.com.br",
                                  "nome": "Drogaria Minas-Brasil",
                                  "busca": "https://www.drogariaminasbrasil.com.br/catalogsearch/result/?q={termo}"},
    "MINAS BRASIL":              {"motor": "auto", "base": "https://www.drogariaminasbrasil.com.br",
                                  "busca": "https://www.drogariaminasbrasil.com.br/catalogsearch/result/?q={termo}"},
    "A NOSSA DROGARIA":          {"motor": "auto", "base": "https://www.anossadrogaria.com.br"},
    "DROGARIAS GLOBO":           {"motor": "auto", "base": "https://www.drogariaglobo.com.br"},
    "FARMA PONTE":               {"motor": "auto", "base": "https://www.farmaponte.com.br"},
    "DROGAO SUPER":              {"motor": "auto", "base": "https://www.drogaosuper.com.br",
                                  "busca": "https://www.drogaosuper.com.br/busca.asp?PalavraChave={termo}"},
    "FARMACIA PERMANENTE":       {"motor": "navegador", "base": "https://www.farmaciapermanente.com.br",
                                  "busca": "https://www.farmaciapermanente.com.br/pesquisa/{termo}"},
    "TAPAJÓS":                   {"motor": "auto", "base": "https://www.santoremedio.com.br", "nome": "Santo Remédio",
                                  "busca": "https://www.santoremedio.com.br/pesquisa?t={termo}"},
    "D 1000":                    {"motor": "vtex", "base": "https://www.drogasmil.com.br", "nome": "Drogasmil",
                                  "bandeiras": [
                                      {"nome": "Drogaria Rosário", "motor": "vtex",
                                       "base": "https://www.drogariarosario.com.br"},
                                      {"nome": "Drogarias Tamoio", "motor": "vtex",
                                       "base": "https://www.drogariastamoio.com.br"},
                                      {"nome": "Farmalife", "motor": "auto",
                                       "base": "https://www.farmalife.com.br"}]},
    # Redepharma tem site institucional (lista lojas), não vende online -> sem preço de site
    "NELFARMA - REDEPHARMA":     {"motor": "auto", "base": None},
    "DROGAVISTA - REDEPHARMA 2": {"motor": "auto", "base": None},
}

CHROME_VERSAO = "140"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      f"(KHTML, like Gecko) Chrome/{CHROME_VERSAO}.0.0.0 Safari/537.36")

PAUSA = 2.0                       # pausa entre página de busca e página de produto (motores HTTP)
TIMEOUT = 25                      # timeout HTTP (s)
TENTATIVAS = 4                    # tentativas HTTP por URL
PRECO_MIN, PRECO_MAX = 0.5, 20000 # faixa de preço aceita (fora disso é leitura errada)
SIMILARIDADE_MIN = 0.6            # quanto o nome no site tem que bater com o SKU para valer como confirmação

RITMO_BASE = {"vtex": 1.5, "sfcc": 4, "nav": 6, "nav_mobile": 5, "nav_visivel": 6}   # s entre consultas por método
RITMO_MAX = 90                    # teto do intervalo após bloqueios
ESFRIAR = 60                      # s de espera base quando o site bloqueia (x nº da tentativa)
MAX_TENTATIVAS = 3                # tentativas por EAN em caso de bloqueio/erro temporário
DESISTIR_APOS = 4                 # bloqueios seguidos para desistir do cliente na rodada
PARALELO_PADRAO = 6               # sites consultados ao mesmo tempo (cada um no seu ritmo; 1 = um de cada vez)
PARCIAL_MIN = 15                  # de quantos em quantos minutos regravar o Excel parcial durante a rodada

CACHE_TTL_HORAS = {               # validade do cache por status (0 = não vale, sempre reconsulta)
    "OK": 24, "SEM ESTOQUE": 6, "SEM PREÇO": 6, "NÃO ENCONTRADO": 8,
    "SITE BLOQUEOU O ROBÔ": 12, "BLOQUEADO": 12,
    # falha da nossa máquina, não do site: assim que o navegador voltar, reconsulta
    "NAVEGADOR INDISPON": 0, "SEM MÉTODO APLICÁVEL": 0,
}

PASTA = Path("preco_site_dados")
ARQ_CACHE = PASTA / "cache_precos.json"
ARQ_METODOS = PASTA / "metodos_clientes.json"
ARQ_HISTORICO = PASTA / "historico_precos.csv"
ARQ_SITES_EXTRA = PASTA / "sites_extra.json"
PASTA_PERFIL = PASTA / "perfil_navegador"
PASTA_LOGS = PASTA / "logs"
PASTA_DEBUG = PASTA / "debug"
ARQ_INDICE = PASTA / "indice_urls.json"
PASTA_SITEMAPS = PASTA / "sitemaps"
ARQ_CACHE_ANTIGO = Path("cache_precos.json")             # versões anteriores gravavam na raiz
ARQ_METODOS_ANTIGO = Path("metodos_clientes_v4.json")

DIAS_SITEMAP = 7                  # de quanto em quanto tempo rebaixar o sitemap de um site
MAX_URLS_SITEMAP = 300000         # teto de URLs guardadas por site

# ================================================================= log ====
log = logging.getLogger("preco_site")


def configura_log(nivel_console=logging.INFO):
    """Console limpo + arquivo detalhado em preco_site_dados/logs/."""
    log.setLevel(logging.DEBUG)
    log.propagate = False
    if log.handlers:
        return
    console = logging.StreamHandler(sys.stdout)
    console.setLevel(nivel_console)
    console.setFormatter(logging.Formatter("%(message)s"))
    log.addHandler(console)
    try:
        PASTA_LOGS.mkdir(parents=True, exist_ok=True)
        arq = logging.FileHandler(PASTA_LOGS / f"preco_site_{datetime.now():%Y%m%d}.log", encoding="utf-8")
        arq.setLevel(logging.DEBUG)
        arq.setFormatter(logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%H:%M:%S"))
        log.addHandler(arq)
    except OSError:
        pass


# =========================================================== utilidades ====
def sem_acento(t) -> str:
    t = unicodedata.normalize("NFKD", "" if t is None else str(t))
    return "".join(c for c in t if not unicodedata.combining(c)).upper().strip()


def dominio(url: Optional[str]) -> str:
    return urlparse(url or "").netloc.lower()


SITES_NORM: dict = {}
SITES_TODOS: dict = {}            # SITES + sites_extra.json, com o nome original (usado no modo régua)
_cfg_resolvido: dict = {}


def _normaliza_cfg(cfg: dict, nome_cliente: str = "") -> dict:
    cfg = dict(cfg or {})
    cfg.setdefault("motor", "auto")
    cfg["base"] = (cfg.get("base") or "").rstrip("/") or None
    if cfg.get("busca"):
        cfg["busca"] = cfg["busca"].replace("{ean}", "{termo}")   # compatível com a v6
    outras = (cfg.get("bandeiras") or cfg.get("alternativos") or [])   # "alternativos" era o nome na v7.0
    cfg["bandeiras"] = [_normaliza_cfg(b) for b in outras if b]
    cfg.pop("alternativos", None)
    if not cfg.get("nome"):
        cfg["nome"] = nome_cliente or (dominio(cfg["base"]).replace("www.", "") if cfg["base"] else "")
    return cfg


def prepara_sites(extra: Optional[dict] = None):
    """Monta o índice de sites (SITES + sites_extra.json) com nomes normalizados."""
    todos = dict(SITES)
    if extra:
        todos.update(extra)
    SITES_NORM.clear()
    SITES_TODOS.clear()
    SITES_TODOS.update(todos)
    _cfg_resolvido.clear()
    for nome, cfg in todos.items():
        SITES_NORM[sem_acento(nome)] = _normaliza_cfg(cfg, nome)


def carrega_sites_extra():
    if not ARQ_SITES_EXTRA.exists():
        prepara_sites()
        return
    try:
        extra = json.loads(ARQ_SITES_EXTRA.read_text("utf-8"))
        if not isinstance(extra, dict):
            raise ValueError("o arquivo precisa ser um objeto {cliente: config}")
        prepara_sites(extra)
        log.info(f"Sites extras carregados de {ARQ_SITES_EXTRA}: {sorted(extra)}")
    except Exception as e:
        log.warning(f"Aviso: não li {ARQ_SITES_EXTRA} ({e}); usando só os sites do código")
        prepara_sites()


def cfg_cliente(cliente) -> Optional[dict]:
    """Config do cliente. Aceita variações do nome ('DROGAL FARMACEUTICA LTDA' -> DROGAL)."""
    nome = sem_acento(cliente)
    if nome in _cfg_resolvido:
        return _cfg_resolvido[nome]
    cfg = SITES_NORM.get(nome)
    if cfg is None and nome:
        candidatos = []
        for chave in SITES_NORM:
            borda = r"(?<![A-Z0-9]){}(?![A-Z0-9])"
            if re.search(borda.format(re.escape(chave)), nome):
                candidatos.append(chave)
            elif len(nome) >= 4 and re.search(borda.format(re.escape(nome)), chave):
                candidatos.append(chave)
        if candidatos:
            chave = max(candidatos, key=len)
            cfg = SITES_NORM[chave]
            log.info(f"Cliente '{cliente}' tratado como '{chave}'")
    _cfg_resolvido[nome] = cfg
    return cfg


def tem_site(cliente) -> bool:
    return bool((cfg_cliente(cliente) or {}).get("base"))


def digito_gtin_valido(gtin: str) -> bool:
    """Valida o dígito verificador de GTIN-8, GTIN-12, GTIN-13 e GTIN-14."""
    if not re.fullmatch(r"\d{8}|\d{12}|\d{13}|\d{14}", gtin or ""):
        return False
    corpo, dv = gtin[:-1], int(gtin[-1])
    soma = sum(int(n) * (3 if i % 2 == 0 else 1) for i, n in enumerate(reversed(corpo)))
    return (10 - soma % 10) % 10 == dv


def normaliza_ean(v) -> Optional[str]:
    """'7896004707037.0', 7.896004707037E+12, ' 7896-0047-07037 ' -> '7896004707037' (ou None se inválido)."""
    if v is None:
        return None
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        s = f"{v:.0f}" if v == int(v) else str(v)
    else:
        s = str(v).strip()
    if re.fullmatch(r"\d+([.,]\d+)?[eE][+-]?\d+", s):
        s = f"{float(s.replace(',', '.')):.0f}"
    s = re.sub(r"[.,]0+$", "", s)
    s = re.sub(r"\D", "", s)
    return s if digito_gtin_valido(s) else None


def canonico_ean(e) -> str:
    return re.sub(r"\D", "", "" if e is None else str(e)).lstrip("0")


def mesmo_ean(a, b) -> bool:
    ca, cb = canonico_ean(a), canonico_ean(b)
    return bool(ca) and ca == cb


def variantes_ean(ean: str) -> list:
    """GTIN original + variações úteis de busca (sem zeros à esquerda / com 13 dígitos)."""
    vals = [ean]
    sem_zeros = ean.lstrip("0")
    if sem_zeros and sem_zeros != ean:
        vals.append(sem_zeros)
    if len(ean) < 13:
        vals.append(ean.zfill(13))
    return list(dict.fromkeys(vals))


LIXO_SKU = re.compile(r"\b(cpr|com|comp|comps|caps?|capsulas?|blt|rev|revest|revb|sol|gotas|cx|x|und|un|"
                      r"a\d+\w*|c\d+\w*)\b")


def termo_busca(sku) -> str:
    """'DIPIRONA 500MG C/30 CPR REV' -> 'dipirona 500mg 30' (termo curto para a caixa de busca)."""
    if sku is None or (isinstance(sku, float) and math.isnan(sku)):
        return ""
    t = sem_acento(sku).lower()
    t = re.sub(r"[()/'\-]", " ", t)
    t = re.sub(r"\bx(\d)", r"\1", t)
    t = LIXO_SKU.sub(" ", t)
    return " ".join([w for w in t.split() if len(w) > 1][:5])


STOP_NOME = {"de", "da", "do", "com", "c", "x", "cx", "und", "un", "para", "e", "a", "o", "em"}


def tokens_nome(t) -> set:
    t = sem_acento(t).lower()
    return {w for w in re.findall(r"[a-z]+|\d+", t) if w not in STOP_NOME and len(w) > 1}


def similaridade_nome(sku, nome_site) -> float:
    """Quanto do SKU (Acelera) aparece no nome do produto no site; números (dose, qtde) pesam mais."""
    a, b = tokens_nome(sku or ""), tokens_nome(nome_site or "")
    if not a or not b:
        return 0.0
    nota = len(a & b) / len(a)
    nums_a, nums_b = {t for t in a if t.isdigit()}, {t for t in b if t.isdigit()}
    if nums_a and not nums_a <= nums_b:   # dose/quantidade diferente -> provável outra apresentação
        nota *= 0.5
    return round(nota, 2)


def num_br(txt) -> Optional[float]:
    """Número em texto humano: 'R$ 1.234,56' -> 1234.56 | '12,9' -> 12.9 | '12.90' -> 12.9 | '1.290' -> 1290."""
    if txt is None:
        return None
    if isinstance(txt, bool):
        return None
    if isinstance(txt, (int, float)):
        return float(txt) if math.isfinite(txt) else None
    m = re.search(r"\d[\d.,]*", str(txt))
    if not m:
        return None
    tok = m.group().rstrip(".,")
    try:
        if "," in tok and "." in tok:
            if tok.rfind(",") > tok.rfind("."):
                tok = tok.replace(".", "").replace(",", ".")
            else:
                tok = tok.replace(",", "")
        elif "," in tok:
            inteiro, _, dec = tok.rpartition(",")
            tok = (inteiro.replace(",", "") + "." + dec) if len(dec) <= 2 else tok.replace(",", "")
        elif "." in tok:
            inteiro, _, dec = tok.rpartition(".")
            tok = tok.replace(".", "") if len(dec) == 3 else (inteiro.replace(".", "") + "." + dec)
        return float(tok)
    except ValueError:
        return None


def num_json(x) -> Optional[float]:
    """Número vindo de JSON/API: ponto decimal tem preferência ('12.90' -> 12.9); cai para num_br."""
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return float(x) if math.isfinite(x) else None
    s = str(x).strip().replace("R$", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return num_br(s)


def preco_valido(v) -> Optional[float]:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(v):
        return None
    return round(v, 2) if PRECO_MIN <= v <= PRECO_MAX else None


RE_PRECO = re.compile(r"R\$\s*\d[\d.]*(?:,\d{1,2})?")
_RE_LIXO = [
    re.compile(r"\d+\s*x\s*(?:de\s*)?(?:sem\s+juros\s*)?(?:R\$)?\s*\d[\d.,]*", re.I),          # 3x de R$ 10,00
    re.compile(r"(?:economi\w*|desconto|cashback|frete|entrega)\W{0,3}(?:de\s*)?R\$\s*\d[\d.,]*", re.I),
    re.compile(r"R\$\s*\d[\d.,]*\s*(?:de\s+)?(?:desconto|economia|cashback|off)", re.I),
    re.compile(r"R\$\s*\d[\d.,]*\s*(?:/|por)\s*(?:un\w*|cada|c[aá]psula|comprimido|ml|g|kg|l)\b", re.I),
    re.compile(r"(?:a partir de|acima de)\s*R\$\s*\d[\d.,]*", re.I),
]


def limpa_texto_preco(t: str) -> str:
    """Tira parcelas, frete, 'economize R$ X', preço por unidade... deixando só preços de venda."""
    for rx in _RE_LIXO:
        t = rx.sub(" ", t)
    return t


def precos_no_texto(texto: str):
    """(preço de venda, preço 'de' riscado) lidos de um trecho de texto. Prefere 'por R$ X'."""
    t = limpa_texto_preco(texto or "")
    todos = [p for p in (preco_valido(num_br(m)) for m in RE_PRECO.findall(t)) if p]
    por = [p for p in (preco_valido(num_br(m.group(1)))
                       for m in re.finditer(r"por\s*:?\s*(R\$\s*\d[\d.]*(?:,\d{1,2})?)", t, re.I)) if p]
    if por:
        menor = min(por)
        return menor, (max(todos) if todos and max(todos) > menor else None)
    if not todos:
        return None, None
    menor = min(todos)
    return menor, (max(todos) if max(todos) > menor else None)


# ============================================================ resultado ====
OK, OK_CONFERIR = "OK", "OK (conferir SKU)"
SEM_ESTOQUE, SEM_PRECO, NAO_ENCONTRADO = "SEM ESTOQUE", "SEM PREÇO", "NÃO ENCONTRADO"
BLOQUEADO, SITE_BLOQUEOU = "BLOQUEADO", "SITE BLOQUEOU O ROBÔ"
NAV_INDISPONIVEL = "NAVEGADOR INDISPONÍVEL"
CAIXA_BUSCA = "CAIXA DE BUSCA NÃO LOCALIZADA"
NAO_E_VTEX = "NÃO É VTEX"
SEM_SITE = "SEM SITE CONFIGURADO"
SEM_METODO = "NENHUM MÉTODO FUNCIONOU"
NAO_CONSULTADO = "NÃO CONSULTADO"

TEMPORARIO = re.compile(r"403|429|HTTP 5|BLOQUEAD|ERRO|TIMEOUT|SITE RESPONDEU|REDE", re.I)
ORDEM_CONFIANCA = {"ALTA": 3, "MÉDIA": 2, "BAIXA": 1, None: 0}


@dataclass
class Resultado:
    status: str
    preco: Optional[float] = None
    preco_de: Optional[float] = None       # preço "de" (riscado / lista), quando o site mostra
    url: Optional[str] = None
    fonte: Optional[str] = None            # vtex | json-ld | meta | dom | texto
    nome_site: Optional[str] = None
    evidencia: Optional[str] = None        # como o EAN foi confirmado
    confianca: Optional[str] = None        # ALTA | MÉDIA | BAIXA
    obs: Optional[str] = None
    metodo: Optional[str] = None
    bandeira: Optional[str] = None         # loja do grupo de onde veio o preço

    @property
    def achou(self) -> bool:
        return self.preco is not None

    @property
    def temporario(self) -> bool:
        return bool(TEMPORARIO.search(self.status or ""))

    def dict(self) -> dict:
        return {k: v for k, v in asdict(self).items() if v is not None}

    @classmethod
    def de_dict(cls, d: dict) -> "Resultado":
        nomes = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in (d or {}).items() if k in nomes})


def _nota_resultado(res: Resultado) -> int:
    """Ordena resultados do melhor para o pior, para escolher o que vale guardar."""
    if res.achou:
        return 5 + ORDEM_CONFIANCA.get(res.confianca, 0)
    if res.status == SEM_ESTOQUE:
        return 4
    if res.status in (NAO_ENCONTRADO, SEM_PRECO):
        return 3                      # resposta definitiva do site: o EAN não está lá
    if res.temporario:
        return 1                      # bloqueio/erro: não sabemos nada
    return 2


def classifica(res: Resultado, nivel: int, evidencia: str, nome_sku, titulo, url) -> Resultado:
    """Define status/confiança do preço lido a partir de como o EAN foi confirmado na página."""
    if res.fonte in ("vtex", "xhr") and (res.evidencia or "").endswith(":ean"):
        nivel, evidencia = 2, res.evidencia      # o próprio JSON da API casou o EAN: evidência máxima
    if nivel >= 2:
        conf, status, ev = "ALTA", OK, evidencia
    elif nivel == 1:
        conf, status, ev = "MÉDIA", OK, evidencia
    else:
        sim = similaridade_nome(nome_sku, titulo) if nome_sku else 0.0
        if sim >= SIMILARIDADE_MIN:
            conf, status, ev = "MÉDIA", OK, f"nome bate {sim:.0%}"
        else:
            conf, status = "BAIXA", OK_CONFERIR
            ev = "EAN não confirmado" + (f" (nome bate {sim:.0%})" if nome_sku else "")
    if res.fonte == "texto" and conf == "ALTA":
        conf, ev = "MÉDIA", ev + ", preço lido do texto"
    if (res.confianca and res.fonte in ("vtex", "xhr")
            and ORDEM_CONFIANCA[res.confianca] < ORDEM_CONFIANCA[conf]):
        conf = res.confianca                     # o motor já rebaixou (ex.: valor da API em centavos)
    if res.status == SEM_ESTOQUE:
        status = SEM_ESTOQUE
    return replace(res, status=status, confianca=conf, evidencia=ev, url=url or res.url,
                   nome_site=res.nome_site or (titulo or None))


# ================================================================ HTTP ====
HEADERS_SESSAO = {
    "User-Agent": UA,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7",
    # Accept-Encoding fica por conta do requests (só anuncia o que consegue descomprimir)
    "Sec-Ch-Ua": f'"Chromium";v="{CHROME_VERSAO}", "Not=A?Brand";v="24", "Google Chrome";v="{CHROME_VERSAO}"',
    "Sec-Ch-Ua-Mobile": "?0",
    "Sec-Ch-Ua-Platform": '"Windows"',
    "Sec-Fetch-Dest": "document",
    "Sec-Fetch-Mode": "navigate",
    "Sec-Fetch-Site": "none",
    "Sec-Fetch-User": "?1",
    "Upgrade-Insecure-Requests": "1",
}
_local = threading.local()


def sessao_http() -> requests.Session:
    """Uma sessão HTTP por thread (requests.Session não é garantida entre threads)."""
    s = getattr(_local, "sessao", None)
    if s is None:
        s = _local.sessao = requests.Session()
        s.headers.update(HEADERS_SESSAO)
    return s


HEADERS_API = {"Accept": "application/json, text/plain, */*", "Sec-Fetch-Dest": "empty",
               "Sec-Fetch-Mode": "cors", "Sec-Fetch-Site": "same-origin"}


class ErroRede(Exception):
    """Falha de rede persistente (DNS, conexão, timeout...)."""


def get(url: str, api: bool = False, **kw) -> requests.Response:
    """GET com retentativas, Retry-After e tolerância a SSL corporativo. Levanta ErroRede se esgotar."""
    headers = dict(HEADERS_API) if api else {}
    if api:
        u = urlparse(url)
        headers["Referer"] = f"{u.scheme}://{u.netloc}/"
    headers.update(kw.pop("headers", None) or {})
    ultimo, verificar = None, True
    for n in range(TENTATIVAS):
        try:
            r = sessao_http().get(url, timeout=TIMEOUT, headers=headers, verify=verificar, **kw)
        except requests.exceptions.SSLError:
            if verificar:
                verificar = False
                log.debug(f"SSL falhou em {dominio(url)}; tentando sem verificar certificado")
                continue
            ultimo = "SSLError"
        except requests.RequestException as e:
            ultimo = type(e).__name__
        else:
            if r.status_code in (408, 425, 429, 500, 502, 503, 504) and n < TENTATIVAS - 1:
                ultimo = f"HTTP {r.status_code}"
                ra = r.headers.get("Retry-After", "")
                espera = min(60, int(ra)) if ra.isdigit() else 3 * (n + 1)
                time.sleep(espera + random.uniform(0.5, 2.0))
                continue
            return r
        time.sleep(2 * (n + 1) + random.uniform(0.5, 2.0))
    raise ErroRede(ultimo or "falha de rede")


RE_TITULO_BLOQ = re.compile(r"just a moment|access denied|attention required|acesso negado|forbidden|"
                            r"request unsuccessful|pardon our interruption|are you a robot|"
                            r"security check|bot detection", re.I)
RE_TEXTO_BLOQ = re.compile(r"verify you are human|verifique se você é humano|access denied|acesso negado|"
                           r"request unsuccessful|_incapsula_resource|unusual traffic|tráfego incomum|"
                           r"cf-chl|challenge-platform|px-captcha|captcha-delivery|errors\.edgesuite\.net|"
                           r"enable javascript and cookies|checking your browser|confirme que você não é um robô",
                           re.I)


def html_bloqueado(html: str) -> bool:
    """Página de desafio/bloqueio (Cloudflare, Akamai, Incapsula, PerimeterX, DataDome...)?"""
    if not html:
        return False
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    if m and RE_TITULO_BLOQ.search(m.group(1)):
        return True
    return len(html) < 25000 and bool(RE_TEXTO_BLOQ.search(html))


# ====================================================== extração HTML ====
def _json_seguro(txt):
    try:
        return json.loads(txt or "")
    except (ValueError, TypeError):
        return None


def _tipo(obj: dict) -> str:
    t = obj.get("@type", "")
    return " ".join(str(x) for x in t) if isinstance(t, list) else str(t)


RE_TIPO_PRODUTO = re.compile(r"Product(?!Group|Collection|Model)")
RE_TIPO_LISTA = re.compile(r"ItemList|SearchResultsPage|CollectionPage")


def _blocos_jsonld(soup) -> list:
    blocos = []
    for tag in soup.find_all("script", type=re.compile(r"ld\+json", re.I)):
        bloco = _json_seguro(tag.string or tag.get_text() or "")
        if bloco is not None:
            blocos.append(bloco)
    return blocos


def _produtos_jsonld(soup) -> list:
    """[(produto, url_do_contexto)] para todo objeto @type Product em qualquer profundidade do JSON-LD."""
    achados, tipos = [], []
    for bloco in _blocos_jsonld(soup):
        pilha, vistos = [(bloco, None)], 0
        while pilha and vistos < 5000:
            obj, url_ctx = pilha.pop()
            vistos += 1
            if isinstance(obj, list):
                pilha.extend((x, url_ctx) for x in obj)
                continue
            if not isinstance(obj, dict):
                continue
            tipo = _tipo(obj)
            tipos.append(tipo)
            if "ListItem" in tipo and isinstance(obj.get("url"), str):
                url_ctx = obj["url"]
            if RE_TIPO_PRODUTO.search(tipo):
                achados.append((obj, url_ctx))
            for v in obj.values():
                if isinstance(v, (dict, list)):
                    pilha.append((v, url_ctx))
    return achados, tipos


def _gtins_jsonld(prod: dict) -> list:
    vals = []
    for campo in ("gtin13", "gtin14", "gtin12", "gtin8", "gtin", "sku", "mpn", "productID"):
        v = prod.get(campo)
        if isinstance(v, (str, int)):
            vals.append(str(v))
    return vals


def _precos_jsonld(prod: dict):
    """(menor preço da oferta, disponível?) de um Product do JSON-LD."""
    ofertas = prod.get("offers") or []
    ofertas = ofertas if isinstance(ofertas, list) else [ofertas]
    precos, disponivel = [], None
    for of in ofertas:
        if not isinstance(of, dict):
            continue
        precos += [num_json(of.get("price")), num_json(of.get("lowPrice"))]
        ps = of.get("priceSpecification")
        for spec in (ps if isinstance(ps, list) else [ps]):
            if isinstance(spec, dict):
                precos.append(num_json(spec.get("price")))
        for sub in of.get("offers") or []:
            if isinstance(sub, dict):
                precos.append(num_json(sub.get("price")))
        disp = str(of.get("availability") or "")
        if disp:
            em_estoque = bool(re.search(r"InStock|LimitedAvailability|OnlineOnly|PreOrder", disp))
            disponivel = em_estoque if disponivel is None else (disponivel or em_estoque)
    precos = [p for p in (preco_valido(x) for x in precos) if p]
    return (min(precos) if precos else None), disponivel


def url_produto_jsonld(html: str, ean: str, base: Optional[str] = None) -> Optional[str]:
    """Numa lista de resultados com JSON-LD, acha a URL do produto cujo GTIN é o EAN procurado."""
    try:
        soup = BeautifulSoup(html or "", "html.parser")
    except Exception:
        return None
    produtos, _ = _produtos_jsonld(soup)
    for prod, url_ctx in produtos:
        if not any(mesmo_ean(g, ean) for g in _gtins_jsonld(prod)):
            continue
        ofertas = prod.get("offers")
        url_oferta = ofertas.get("url") if isinstance(ofertas, dict) else None
        for u in (prod.get("url"), url_ctx, url_oferta, prod.get("@id")):
            if isinstance(u, str) and u.startswith(("http", "/")):
                return urljoin(base or "", u) if base else u
    return None


SELETORES_META = [
    ('[data-price-type="finalPrice"]', "data-price-amount"),
    ('meta[property="product:price:amount"]', "content"),
    ('meta[property="og:price:amount"]', "content"),
    ('meta[itemprop="price"]', "content"),
    ('[itemprop="price"]', "content"),
    (".sales .value", "content"),
]


def extrai_preco_html(html: str, ean: Optional[str] = None) -> Optional[Resultado]:
    """Preço de uma página de produto: JSON-LD (preferindo o Product com o GTIN certo) > meta/microdata > texto."""
    try:
        soup = BeautifulSoup(html or "", "html.parser")
    except Exception:
        return None
    h1 = soup.find("h1")
    titulo = (h1.get_text(" ", strip=True) if h1 else "") or None

    produtos, _ = _produtos_jsonld(soup)
    if produtos:
        batem = [p for p, _ in produtos if ean and any(mesmo_ean(g, ean) for g in _gtins_jsonld(p))]
        ordem = batem or ([produtos[0][0]] if len(produtos) == 1 else [])
        for prod in ordem:
            preco, disponivel = _precos_jsonld(prod)
            if preco:
                nome = prod.get("name")
                return Resultado(SEM_ESTOQUE if disponivel is False else OK, preco=preco, fonte="json-ld",
                                 nome_site=str(nome) if nome else titulo,
                                 evidencia="json-ld:gtin" if prod in batem else None)

    for sel, attr in SELETORES_META:
        el = soup.select_one(sel)
        if el is not None:
            preco = preco_valido(num_json(el.get(attr)) if el.get(attr) else num_br(el.get_text()))
            if preco:
                return Resultado(OK, preco=preco, fonte="meta", nome_site=titulo)

    if h1:
        bloco = h1.find_parent(["main", "section", "article"]) or h1.parent or soup
        preco, de = precos_no_texto(bloco.get_text(" "))
        if preco:
            return Resultado(OK, preco=preco, preco_de=de, fonte="texto", nome_site=titulo)
    return None


def confirma_ean_pagina(html: str, ean: str):
    """(nível, evidência): 2 = GTIN estruturado (json-ld/microdata/estado JS) | 1 = texto visível | 0 = nada."""
    try:
        soup = BeautifulSoup(html or "", "html.parser")
    except Exception:
        return 0, "html inválido"
    produtos, _ = _produtos_jsonld(soup)
    for prod, _ in produtos:
        if any(mesmo_ean(g, ean) for g in _gtins_jsonld(prod)):
            return 2, "json-ld"
    for el in soup.select('[itemprop^="gtin"], [itemprop="sku"], [itemprop="productID"], '
                          'meta[property*="gtin" i], meta[name*="gtin" i], meta[property*="ean" i], meta[name*="ean" i]'):
        if mesmo_ean(el.get("content") or el.get_text(), ean):
            return 2, "microdata"
    valor = "0*" + re.escape(canonico_ean(ean))                 # aceita zeros à esquerda (07896..., 00007896...)
    scripts = " ".join(t.get_text(" ") for t in soup.find_all("script"))
    if re.search(rf'(?i)(gtin\w*|ean\w*|barcode|codigo_?barras?|cod_?barras?)\\?["\']?\s*[:=]\s*\\?["\']?\s*'
                 rf'{valor}(?!\d)', scripts):
        return 2, "estado estruturado"
    texto = soup.get_text(" ")
    if re.search(rf"(?i)(ean|gtin|c[óo]digo de barras|cod\.?\s*barras)\D{{0,25}}{valor}(?!\d)", texto):
        return 1, "texto da página"
    return 0, "não confirmado"


def pagina_de_produto(url: str, html: str, ean: str) -> bool:
    """Heurística: a página atual é a página de UM produto (e não uma lista de resultados)?"""
    soup = BeautifulSoup(html or "", "html.parser")
    produtos, tipos = _produtos_jsonld(soup)
    tem_lista = any(RE_TIPO_LISTA.search(t) for t in tipos)
    if produtos:
        return not tem_lista
    caminho = urlparse(url or "").path.lower()
    if re.search(r"\.html?$|/p/?$|/produto/|/product/", caminho) and not RE_URL_BUSCA.search(url or ""):
        return True
    return False


RE_URL_BUSCA = re.compile(r"busca|search|pesquis|catalogsearch|[?&](q|w|s|termo\w*|query)=", re.I)


def monta_url_busca(modelo: str, termo: str) -> str:
    """Preenche o {termo} no modelo de URL de busca.

    Query string ('?q={termo}') aceita '+' no lugar do espaço; caminho ('/pesquisa/{termo}') não —
    ali o espaço tem que virar %20, senão o site procura por um termo com '+' no meio e não acha nada.
    """
    antes = modelo.split("{termo}")[0]
    codificado = quote_plus(termo) if "?" in antes else quote(termo, safe="")
    return modelo.replace("{termo}", codificado)


# ================================== JSON das APIs que a própria página chama ==
# O site (ainda mais na versão mobile) busca o produto por JSON. Interceptar essa
# resposta dá preço + EAN na mesma estrutura: é a evidência mais forte que existe
# e não quebra quando o layout muda.
RE_CHAVE_EAN = re.compile(r"^(ean\w*|gtin\w*|barcode|barras|codigo_?(de_?)?barras?|cod_?(igo)?_?barras?|"
                          r"referenceid|refid|alternateids?_?ean)$", re.I)
RE_CHAVE_IGNORAR = re.compile(r"installment|parcel|juros|frete|shipping|entrega|cashback|desconto|discount|"
                              r"tax|imposto|similar|relacionad|recomend|suggest", re.I)
RE_CHAVE_PRECO = re.compile(r"^(price|preco|pre[cç]o|sellingprice|bestprice|finalprice|saleprice|currentprice|"
                            r"pricewithdiscount|valor|valorvenda|precovenda|preco_?por|precopor|pricevalue|"
                            r"spotprice|unitmultiplierprice|lowprice)$", re.I)
RE_CHAVE_PRECO_DE = re.compile(r"^(listprice|oldprice|originalprice|regularprice|precode|preco_?de|precooriginal|"
                               r"pricewithoutdiscount|valordelista|maxprice|highprice|precotabela)$", re.I)
RE_CHAVE_NOME = re.compile(r"^(name|nome|productname|nomecompleto|namecomplete|title|titulo|descricao|"
                           r"description|shortdescription|nome_?produto)$", re.I)
RE_CHAVE_ESTOQUE = re.compile(r"^(available|isavailable|disponivel|instock|emestoque|availablequantity|"
                              r"quantity|estoque|qtdestoque|saleable)$", re.I)
RE_CHAVE_URL = re.compile(r"^(url|link|linktext|permalink|detailurl|producturl|slug)$", re.I)


def _dicts_json(obj, limite: int = 20000):
    """Todos os dicionários do JSON, em qualquer profundidade (com teto, para não travar)."""
    pilha, vistos = [obj], 0
    while pilha and vistos < limite:
        atual = pilha.pop()
        vistos += 1
        if isinstance(atual, dict):
            yield atual
            pilha.extend(v for v in atual.values() if isinstance(v, (dict, list)))
        elif isinstance(atual, list):
            pilha.extend(v for v in atual if isinstance(v, (dict, list)))


def _valores_rasos(d: dict):
    """Valores escalares do dicionário + os de listas simples e de pares {Key/Value}."""
    for v in d.values():
        if isinstance(v, (str, int, float)):
            yield v
        elif isinstance(v, list):
            for x in v:
                if isinstance(x, (str, int, float)):
                    yield x
                elif isinstance(x, dict):
                    for chave in ("Value", "value", "valor"):
                        if isinstance(x.get(chave), (str, int, float)):
                            yield x[chave]


def _identifica_ean(d: dict, ean: str) -> bool:
    """O dicionário diz, ele mesmo, que é o produto deste EAN?"""
    for chave, valor in d.items():
        if not RE_CHAVE_EAN.search(str(chave)):
            continue
        if isinstance(valor, (str, int)):
            if mesmo_ean(valor, ean):
                return True
        elif isinstance(valor, list):
            for x in valor:
                if isinstance(x, (str, int)) and mesmo_ean(x, ean):
                    return True
                if isinstance(x, dict) and any(mesmo_ean(x.get(k), ean) for k in ("Value", "value", "valor")):
                    return True
    return False


def _candidatos_preco(d: dict, profundidade: int = 4):
    """[(valor cru, é preço 'de'?)] do dicionário e de seus filhos (para pegar sellers/offers aninhados)."""
    achados, pilha = [], [(d, 0)]
    while pilha:
        atual, nivel = pilha.pop()
        if not isinstance(atual, dict) or nivel > profundidade:
            continue
        for chave, valor in atual.items():
            if RE_CHAVE_IGNORAR.search(str(chave)):     # parcelas, frete, produtos relacionados
                continue
            if isinstance(valor, (int, float, str)) and not isinstance(valor, bool):
                if RE_CHAVE_PRECO.search(str(chave)):
                    achados.append((valor, False))
                elif RE_CHAVE_PRECO_DE.search(str(chave)):
                    achados.append((valor, True))
            elif isinstance(valor, dict):
                pilha.append((valor, nivel + 1))
            elif isinstance(valor, list):
                pilha.extend((x, nivel + 1) for x in valor if isinstance(x, dict))
    return achados


FAIXA_CENTAVOS = (1000, 200000)      # 1990 -> 19,90; acima disso não dá para supor centavos com segurança


def _preco_api(valor, referencia: Optional[float] = None):
    """Normaliza preço vindo de API. Devolve (preço, é_suposição) — algumas APIs mandam 1990 por 19,90."""
    v = num_json(valor)
    if v is None:
        return None, False
    direto, centavos = preco_valido(v), preco_valido(v / 100)
    inteiro = float(v).is_integer()
    pode_centavos = bool(centavos) and inteiro and FAIXA_CENTAVOS[0] <= v <= FAIXA_CENTAVOS[1]
    if referencia:                                        # temos um preço lido da página: escolhe o mais próximo
        opcoes = [(abs(direto - referencia), direto)] if direto else []
        if pode_centavos:
            opcoes.append((abs(centavos - referencia), centavos))
        if not opcoes:
            return None, False
        erro, preco = min(opcoes)
        return preco, erro > max(0.05, 0.02 * referencia)   # bateu com a página: não é suposição
    if direto and not pode_centavos:
        return direto, False
    if pode_centavos:
        return centavos, True
    return direto, False


def valor_aparece_no_texto(preco: Optional[float], texto: Optional[str]) -> bool:
    """O preço, escrito como o brasileiro escreve (19,90 / 1.234,56), aparece no texto da página?"""
    if not preco or not texto:
        return False
    inteiro, _, dec = f"{preco:.2f}".partition(".")
    if f"{inteiro},{dec}" in texto:
        return True
    return len(inteiro) > 3 and f"{int(inteiro):,}".replace(",", ".") + f",{dec}" in texto


def preco_em_json(dados, ean: str, referencia: Optional[float] = None,
                  texto: Optional[str] = None) -> Optional[Resultado]:
    """Acha, no JSON de uma API, o objeto que declara este EAN e lê o preço dele.

    referencia = preço já lido da página (resolve API em centavos);
    texto      = texto visível da página (confirma o valor e devolve a confiança ALTA).
    """
    for d in _dicts_json(dados):
        if not _identifica_ean(d, ean):
            continue
        brutos = _candidatos_preco(d)
        vendas = [(p, sup) for p, sup in (_preco_api(v, referencia) for v, e_de in brutos if not e_de) if p]
        if not vendas:                                    # sem preço de venda (só "de", ou nada legível)
            continue
        preco, assumiu = min(vendas)                      # o menor é o preço praticado; a suposição é a dele
        if assumiu and valor_aparece_no_texto(preco, texto):
            assumiu = False                               # o valor convertido está escrito na página: confirmado
        # o preço "de" se confere contra o próprio preço de venda, não contra a página
        de = max([p for p, _ in (_preco_api(v, preco) for v, e_de in brutos if e_de) if p] or [0]) or None
        nome = next((str(v) for k, v in d.items()
                     if RE_CHAVE_NOME.search(str(k)) and isinstance(v, str) and len(v) > 3), None)
        # a própria resposta costuma trazer o link da página do produto: guardamos para ir direto na próxima vez
        url = next((v for k, v in d.items()
                    if RE_CHAVE_URL.search(str(k)) and isinstance(v, str) and v.startswith(("/", "http"))), None)
        disponivel = None
        for k, v in d.items():
            if not RE_CHAVE_ESTOQUE.search(str(k)):
                continue
            if isinstance(v, bool):
                disponivel = v if disponivel is None else (disponivel or v)
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                disponivel = (v > 0) if disponivel is None else (disponivel or v > 0)
        obs = "valor da API interpretado como centavos" if assumiu else None
        return Resultado(SEM_ESTOQUE if disponivel is False else OK, preco=preco,
                         preco_de=de if de and de > preco else None, url=url, fonte="xhr", nome_site=nome,
                         evidencia="api da página:ean", confianca="MÉDIA" if assumiu else "ALTA", obs=obs)
    return None


def preco_nas_capturas(payloads: list, ean: str, referencia: Optional[float] = None,
                       texto: Optional[str] = None):
    """Varre as respostas JSON capturadas. Devolve (Resultado, url da API) ou (None, None)."""
    for url, dados in payloads or []:
        try:
            res = preco_em_json(dados, ean, referencia, texto)
        except Exception:
            continue
        if res is not None:
            log.debug(f"preço veio do JSON da página: {url[:120]}")
            return res, url
    return None, None


# ================================================================ VTEX ====
def _item_bate_ean(item: dict, ean: str) -> bool:
    if mesmo_ean(item.get("ean"), ean):
        return True
    for ref in item.get("referenceId") or []:
        if isinstance(ref, dict) and mesmo_ean(ref.get("Value"), ean):
            return True
    return False


def busca_vtex(cfg: dict, ean: str, nome=None) -> Resultado:
    """API pública de catálogo da VTEX. Filtra pelo EAN (com variações), depois RefId, depois texto livre."""
    base = cfg["base"]
    consultas = [f"fq=alternateIds_Ean:{v}" for v in variantes_ean(ean)]
    consultas += [f"fq=alternateIds_RefId:{ean}", f"ft={ean}"]
    dados, ultimo = None, None
    for q in consultas:
        r = get(f"{base}/api/catalog_system/pub/products/search?{q}&_from=0&_to=9", api=True)
        if r.status_code in (401, 403, 429):
            return Resultado(f"HTTP {r.status_code}")
        if r.status_code not in (200, 206):
            ultimo = r.status_code
            continue
        try:
            lista = r.json()
        except ValueError:
            return Resultado(NAO_E_VTEX)
        if not isinstance(lista, list):
            return Resultado(NAO_E_VTEX)
        dados = lista
        if lista:
            break
        time.sleep(random.uniform(0.3, 0.8))
    if dados is None:
        return Resultado(NAO_E_VTEX if ultimo in (None, 404) else f"HTTP {ultimo}")
    if not dados:
        return Resultado(NAO_ENCONTRADO)

    achou_item, sem_estoque, indisponivel = False, None, None
    for prod in dados:
        if not isinstance(prod, dict):
            continue
        link_text = prod.get("linkText")
        url = f"{base}/{link_text}/p" if link_text else prod.get("link")
        for item in prod.get("items") or []:
            if not isinstance(item, dict) or not _item_bate_ean(item, ean):
                continue
            achou_item = True
            nome_site = prod.get("productName") or item.get("nameComplete") or item.get("name")
            for seller in item.get("sellers") or []:
                of = (seller or {}).get("commertialOffer") or {}
                disponivel = (num_json(of.get("AvailableQuantity")) or 0) > 0 or bool(of.get("IsAvailable"))
                preco = preco_valido(of.get("Price"))
                de = preco_valido(of.get("ListPrice"))
                if not preco:
                    if not disponivel and indisponivel is None:      # esgotado: a VTEX zera o Price
                        indisponivel = Resultado(SEM_ESTOQUE, preco_de=de, url=url, fonte="vtex",
                                                 nome_site=nome_site, evidencia="api:ean", confianca="ALTA",
                                                 obs="produto esgotado no site (sem preço de venda)")
                    continue
                res = Resultado(OK, preco=preco, preco_de=de if de and de > preco else None, url=url,
                                fonte="vtex", nome_site=nome_site, evidencia="api:ean", confianca="ALTA")
                if disponivel:
                    return res
                sem_estoque = sem_estoque or replace(res, status=SEM_ESTOQUE)
    if sem_estoque:
        return sem_estoque
    if indisponivel:
        return indisponivel
    return Resultado(SEM_PRECO if achou_item else NAO_ENCONTRADO)


# ==================================================== SFCC (Araujo) ======
def _link_produto_sfcc(soup, ean: str, base: str) -> Optional[str]:
    """Na lista de resultados do SFCC: tile com data-pid = EAN, imagem com o EAN no src, ou 1º tile."""
    for tile in soup.select("[data-pid]"):
        if mesmo_ean(tile.get("data-pid"), ean):
            a = tile.select_one(".pdp-link a[href], a.link[href], a[href]")
            if a:
                return urljoin(base, a["href"])
    for img in soup.find_all("img"):
        src = " ".join(str(img.get(a, "")) for a in ("src", "data-src", "srcset", "data-srcset"))
        if ean in src or (len(canonico_ean(ean)) >= 8 and canonico_ean(ean) in src):
            a = img.find_parent("a", href=True)
            if a:
                return urljoin(base, a["href"])
    return None


def busca_sfcc(cfg: dict, ean: str, nome=None, indice=None) -> Resultado:
    """Salesforce Commerce Cloud: busca por URL; se cair direto na página do produto, lê; senão acha o link."""
    def le_pagina(html, url):
        res = extrai_preco_html(html, ean)
        if not res:
            return None
        nivel, evid = confirma_ean_pagina(html, ean)
        return classifica(res, nivel, evid, nome, res.nome_site, url)

    conhecida = indice.url(cfg.get("base"), ean) if indice is not None else None
    if conhecida:                                   # já sabemos a página deste produto: vai direto
        try:
            rc = get(conhecida)
            if rc.status_code == 200 and not html_bloqueado(rc.text) and pagina_de_produto(rc.url, rc.text, ean):
                pronto = le_pagina(rc.text, rc.url)
                if pronto is not None and pronto.confianca in ("ALTA", "MÉDIA"):
                    return replace(pronto, obs="página já conhecida")
        except ErroRede:
            pass
        indice.esquece(cfg.get("base"), ean)

    r = get(monta_url_busca(cfg["busca"], ean))
    if r.status_code != 200:
        return Resultado(f"HTTP {r.status_code}")
    if html_bloqueado(r.text):
        return Resultado(BLOQUEADO)

    def le_pdp(html, url):
        return le_pagina(html, url) or Resultado(SEM_PRECO, url=url)

    if pagina_de_produto(r.url, r.text, ean):
        return le_pdp(r.text, r.url)
    soup = BeautifulSoup(r.text, "html.parser")
    url_prod = url_produto_jsonld(r.text, ean, cfg["base"]) or _link_produto_sfcc(soup, ean, cfg["base"])
    if not url_prod:
        return Resultado(NAO_ENCONTRADO)
    time.sleep(PAUSA + random.uniform(0.5, 1.5))
    r2 = get(url_prod)
    if r2.status_code != 200:
        return Resultado(f"HTTP {r2.status_code}")
    if html_bloqueado(r2.text):
        return Resultado(BLOQUEADO)
    return le_pdp(r2.text, r2.url)


# ======================================================= NAVEGADOR =======
JS_INIT = r"""
try { if (navigator.webdriver) Object.defineProperty(navigator, 'webdriver', {get: () => false}); } catch (e) {}
try { Object.defineProperty(navigator, 'languages', {get: () => ['pt-BR', 'pt', 'en-US', 'en']}); } catch (e) {}
"""

# Acha o link do produto numa página de resultados. Ordem: href com o EAN > imagem com o EAN >
# seletores conhecidos de card > card genérico com imagem e "R$". Devolve null se a página diz
# que não achou nada (evita pegar "produtos recomendados").
JS_PRODUTO = r"""(ean) => {
    const aqui = location.origin + location.pathname;
    const ruim = /\/(busca|search|catalogsearch|pesquisa|carrinho|cart|login|conta|minha-conta|account|checkout|categoria|categorias|departamento|departamentos|marca|marcas|institucional|blog|ajuda|faq|sobre)(\/|$|\?)/i;
    const semResultado = /(nenhum (resultado|produto encontrado|item encontrado)|n[aã]o (encontramos|foram encontrados|h[aá] resultados|foi encontrado)|\b0 resultados?\b|sem resultados?|nada encontrado|no results)/i;
    const conteudo = document.querySelector('main') || document.body;
    if (conteudo && semResultado.test((conteudo.innerText || '').slice(0, 4000))) return null;
    const valido = a => {
        try {
            const href = a.getAttribute('href') || '';
            if (/^(javascript:|mailto:|tel:|#)/i.test(href)) return false;
            const u = new URL(a.href, location.href);
            if (u.host !== location.host) return false;
            if ((u.origin + u.pathname) === aqui) return false;
            if (ruim.test(u.pathname)) return false;
            return u.pathname.length > 3;
        } catch (e) { return false; }
    };
    const foraSempre = 'header, footer, nav';
    const foraEstrito = foraSempre + ', [class*="carousel" i], [class*="slider" i], [class*="swiper" i], [class*="recomend" i], [class*="related" i], [class*="vitrine" i], [class*="shelf" i], [class*="sugest" i]';
    const ean2 = ean.replace(/^0+/, '');
    const temEan = s => !!s && (s.includes(ean) || (ean2.length >= 8 && s.includes(ean2)));
    const procura = (fora) => {
        const dentro = a => !a.closest(fora);
        for (const a of document.querySelectorAll('a[href]')) {
            if (!valido(a) || !dentro(a)) continue;
            try { if (temEan(new URL(a.href, location.href).pathname)) return a.href; } catch (e) {}
        }
        for (const img of document.querySelectorAll('img')) {
            const src = (img.currentSrc || '') + ' ' + (img.src || '') + ' ' + (img.srcset || '') + ' ' + (img.dataset.src || '');
            const a = img.closest('a[href]');
            if (a && valido(a) && dentro(a) && temEan(src)) return a.href;
        }
        const seletores = ['a.product-item-link', '.pdp-link a[href]', 'a.product-name', '[class*="product-summary" i] a[href]',
                           '[class*="productCard" i] a[href]', '[class*="product-card" i] a[href]', '[class*="produto-card" i] a[href]',
                           '[data-testid*="product" i] a[href]', 'li.product a[href]', '.product a[href]', '.produto a[href]',
                           '.prateleira a[href]', '[class*="gallery" i] a[href]'];
        for (const sel of seletores) {
            for (const a of document.querySelectorAll(sel)) if (valido(a) && dentro(a)) return a.href;
        }
        const genericos = ['main a[href]', '[class*="result" i] a[href]', '[class*="search" i] a[href]', '[class*="list" i] a[href]', 'a[href]'];
        for (const sel of genericos) {
            for (const a of document.querySelectorAll(sel)) {
                if (!valido(a) || !dentro(a)) continue;
                const card = a.closest('li, article, [class*="product" i], [class*="produto" i], [class*="card" i], [class*="item" i], [class*="tile" i]') || a;
                if ((a.querySelector('img') || card.querySelector('img')) && /R\$\s*\d/.test(card.innerText || '')) return a.href;
            }
        }
        return null;
    };
    return procura(foraEstrito) || procura(foraSempre);   // 2ª passada: sites que chamam a lista de "vitrine"/"shelf"
}"""

# Candidatos a preço no DOM (elementos com "price"/"preco" no nome), perto do h1, ignorando vitrines,
# parcelas e frete. Vem ordenado: não riscado primeiro, fonte maior primeiro.
JS_PRECO_DOM = r"""() => {
    const h1 = document.querySelector('h1');
    const fora = '[class*="shelf" i], [class*="carousel" i], [class*="related" i], [class*="recomend" i], [class*="slider" i], [class*="swiper" i], [class*="installment" i], [class*="parcel" i], [class*="shipping" i], [class*="frete" i], [class*="minicart" i], [class*="vitrine" i], header, footer, nav';
    const lixo = /(\d+\s*x\b|parcela|frete|entrega|economi|cashback|por unidade|\/un|cada|c[aá]psula|comprimido|a partir de|desconto)/i;
    const topoH1 = h1 ? h1.getBoundingClientRect().top + window.scrollY : 0;
    const cands = [];
    for (const e of document.querySelectorAll('[class*="price" i], [class*="preco" i], [id*="price" i], [data-testid*="price" i], [data-price-amount]')) {
        if (e.offsetParent === null || e.closest(fora)) continue;
        const t = (e.innerText || '').trim();
        if (!/R\$\s*\d/.test(t) || t.length > 80 || lixo.test(t)) continue;
        if (e.querySelector('[class*="price" i], [class*="preco" i]')) continue;
        const r = e.getBoundingClientRect();
        const topo = r.top + window.scrollY;
        if (h1 && (topo < topoH1 - 150 || topo > topoH1 + 1200)) continue;
        const riscado = getComputedStyle(e).textDecorationLine.includes('line-through')
            || !!e.closest('del, s, strike, [class*="old" i], [class*="list-price" i], [class*="listprice" i], [class*="price-de" i], [class*="preco-de" i], [class*="precode" i], [class*="strike" i], [class*="riscado" i], [class*="regular-price" i]');
        cands.push({t, fonte: parseFloat(getComputedStyle(e).fontSize) || 0, riscado});
    }
    cands.sort((x, y) => (x.riscado - y.riscado) || (y.fonte - x.fonte));
    return cands.slice(0, 5);
}"""

JS_INFO_PAGINA = r"""() => {
    const h1 = document.querySelector('h1');
    const compra = document.querySelectorAll('button[class*="buy" i], button[class*="comprar" i], button[class*="add-to-cart" i], button[class*="addtocart" i], [class*="add-to-cart" i] button, [data-testid*="add-to-cart" i], [data-testid*="buy" i], button[id*="buy" i], button[id*="comprar" i], form[action*="cart" i] button[type="submit"]').length;
    return {h1: h1 ? (h1.innerText || '').trim() : '', titulo: document.title || '', compra, url: location.href};
}"""

SELETORES_BUSCA = ['input[type="search"]', 'input[name="q"]', 'input[name="w"]', 'input[name="s"]',
                   'input[name*="busca" i]', 'input[name*="search" i]', 'input[name*="termo" i]',
                   'input[id*="search" i]', 'input[id*="busca" i]',
                   'input[placeholder*="usca" i]', 'input[placeholder*="rocur" i]',
                   'input[placeholder*="esquis" i]', 'input[placeholder*="que voc" i]',
                   'input[aria-label*="usca" i]', 'input[aria-label*="esquis" i]']

# No celular a caixa de busca costuma ficar escondida atrás de uma lupa: clicar nisso primeiro.
BOTOES_ABRIR_BUSCA = ['[aria-label*="usca" i]', '[aria-label*="esquis" i]', '[aria-label*="search" i]',
                      'button[class*="search" i]', 'button[class*="busca" i]', 'button[class*="lupa" i]',
                      '[data-testid*="search" i]', 'a[href*="/busca" i]', 'a[href*="/search" i]',
                      'svg[class*="search" i]', 'i[class*="search" i]']

BOTOES_POPUP = ['#onetrust-accept-btn-handler', 'button:has-text("Aceitar")', 'button:has-text("Aceito")',
                'button:has-text("Concordo")', 'button:has-text("Entendi")', 'button:has-text("Prosseguir")',
                '[aria-label*="echar" i]', '[aria-label*="close" i]',
                'button:has-text("Agora não")', 'button:has-text("Depois")', 'button:has-text("Não, obrigado")',
                # interstícios de "baixe o app" que só aparecem no celular
                'button:has-text("Continuar no site")', 'a:has-text("Continuar no site")',
                'button:has-text("Continuar no navegador")', 'button:has-text("Ficar no site")',
                'button:has-text("Agora não, obrigado")', '[class*="smartbanner" i] [class*="close" i]',
                '[id*="smartbanner" i] button', '[class*="app-banner" i] button[class*="close" i]']

# ------------------------------------------------------------------ mobile --
# Perfis coerentes: UA, viewport, touch e engine combinando entre si.
# (UA de iPhone em cima do Chromium é fácil de detectar - por isso o iPhone só vai no WebKit.)
ANDROID_VERSAO = "14"
UA_ANDROID = (f"Mozilla/5.0 (Linux; Android {ANDROID_VERSAO}; Pixel 8) AppleWebKit/537.36 "
              f"(KHTML, like Gecko) Chrome/{CHROME_VERSAO}.0.0.0 Mobile Safari/537.36")
UA_IPHONE = ("Mozilla/5.0 (iPhone; CPU iPhone OS 17_5 like Mac OS X) AppleWebKit/605.1.15 "
             "(KHTML, like Gecko) Version/17.5 Mobile/15E148 Safari/604.1")

APARELHO_ANDROID = {
    "user_agent": UA_ANDROID,
    "viewport": {"width": 412, "height": 915},
    "device_scale_factor": 2.625,
    "is_mobile": True,
    "has_touch": True,
    "extra_http_headers": {"Sec-Ch-Ua-Mobile": "?1", "Sec-Ch-Ua-Platform": '"Android"'},
}
APARELHO_IPHONE = {
    "user_agent": UA_IPHONE,
    "viewport": {"width": 393, "height": 852},
    "device_scale_factor": 3,
    "is_mobile": True,
    "has_touch": True,
}

# Mantém navigator.platform / userAgentData coerentes com o UA de Android (o Chromium
# roda em Windows/Linux e entregaria "Win32" com UA de celular - incoerência fácil de pegar).
JS_INIT_ANDROID = r"""
try { Object.defineProperty(navigator, 'platform', {get: () => 'Linux armv81'}); } catch (e) {}
try {
    if (navigator.userAgentData) {
        const marcas = [{brand: 'Chromium', version: '%(v)s'}, {brand: 'Google Chrome', version: '%(v)s'}, {brand: 'Not=A?Brand', version: '24'}];
        Object.defineProperty(navigator.userAgentData, 'mobile', {get: () => true});
        Object.defineProperty(navigator.userAgentData, 'platform', {get: () => 'Android'});
        Object.defineProperty(navigator.userAgentData, 'brands', {get: () => marcas});
        const alta = navigator.userAgentData.getHighEntropyValues.bind(navigator.userAgentData);
        navigator.userAgentData.getHighEntropyValues = async (h) => Object.assign(await alta(h), {
            platform: 'Android', platformVersion: '%(a)s', mobile: true, model: 'Pixel 8',
            architecture: '', bitness: '', brands: marcas});
    }
} catch (e) {}
try { Object.defineProperty(navigator, 'maxTouchPoints', {get: () => 5}); } catch (e) {}
try { Object.defineProperty(navigator, 'hardwareConcurrency', {get: () => 8}); } catch (e) {}
""" % {"v": CHROME_VERSAO, "a": ANDROID_VERSAO + ".0.0"}


class NavegadorIndisponivel(Exception):
    pass


# erros que valem para qualquer perfil de navegador (não adianta tentar celular/visível depois)
ERRO_INSTALACAO = re.compile(r"Executable doesn't exist|playwright install|missing dependencies|"
                             r"Host system is missing", re.I)


def checa_navegador() -> Optional[str]:
    """Confere, sem abrir nada, se o Playwright e o Chromium estão instalados. None = tudo certo."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return "falta o playwright: pip install playwright"
    try:
        with sync_playwright() as pw:
            caminho = pw.chromium.executable_path
    except Exception as e:
        return f"{type(e).__name__}: {str(e).splitlines()[0][:150]}"
    if caminho and not Path(caminho).exists():
        return "falta o Chromium: python -m playwright install chromium"
    return None


def tem_tela() -> bool:
    """Existe uma tela para abrir janela de navegador? (Colab, servidor e cron não têm)"""
    if sys.platform.startswith("win") or sys.platform == "darwin":
        return True
    return bool(os.environ.get("DISPLAY") or os.environ.get("WAYLAND_DISPLAY"))


class NavegadorMorreu(Exception):
    """O processo do navegador caiu; é preciso abrir outro."""


class Navegador:
    """Chromium via Playwright com perfil persistente (cookies), comportamento humano e detecção de bloqueio.

    mobile=True emula um Pixel 8 (Chromium) — UA, viewport, touch e client hints coerentes entre si —
    e cai para WebKit/iPhone se o Chromium não abrir. O site mobile costuma ter HTML mais simples e
    buscar preço por JSON, que o robô intercepta (ver captura de XHR).
    """

    def __init__(self, visivel: bool = False, busca_aprendida: Optional[dict] = None, mobile: bool = False,
                 indice: Optional["IndiceUrls"] = None, sufixo_perfil: str = ""):
        try:
            from playwright.sync_api import Error as ErroPW, TimeoutError as TimeoutPW, sync_playwright
        except ImportError:
            raise NavegadorIndisponivel("pip install playwright  &&  python -m playwright install chromium")
        self.ErroPW, self.TimeoutPW = ErroPW, TimeoutPW
        self.visivel = visivel
        self.mobile = mobile
        self.indice = indice
        self.sufixo_perfil = sufixo_perfil     # cada trabalhador em paralelo precisa do seu perfil (o Chromium trava o dele)
        self.aparelho = "desktop"
        self.aquecidos: set = set()
        self.busca_aprendida = busca_aprendida if busca_aprendida is not None else {}
        self.navegacoes = 0
        self.url_digitada: Optional[str] = None
        self.respostas: list = []          # respostas JSON vistas na página (captura de XHR)
        self.ctx = None
        self._pw = sync_playwright().start()
        try:
            self.ctx = self._lancar()
            self.page = self._pagina()
            ua = self.page.evaluate("() => navigator.userAgent") or ""
            if "Headless" in ua:                       # o headless denuncia "HeadlessChrome": corrige mantendo a versão real
                self.ctx.close()
                self.ctx = self._lancar(user_agent=ua.replace("HeadlessChrome", "Chrome"))
                self.page = self._pagina()
            self._liga_captura()
            log.debug(f"navegador pronto ({self.aparelho}): {self.page.evaluate('() => navigator.userAgent')}")
        except NavegadorIndisponivel:
            self._pw.stop()
            raise
        except Exception:
            self.fechar()
            raise

    @property
    def perfil(self) -> Path:
        """Perfil separado por aparelho: cookies de celular e de desktop não se misturam."""
        return PASTA_PERFIL.with_name(PASTA_PERFIL.name + self.sufixo_perfil + ("_mobile" if self.mobile else ""))

    def _lancar(self, **extra):
        self.perfil.parent.mkdir(parents=True, exist_ok=True)
        base = dict(headless=not self.visivel, locale="pt-BR", timezone_id="America/Sao_Paulo",
                    ignore_https_errors=True,
                    args=["--disable-blink-features=AutomationControlled", "--no-sandbox",
                          "--disable-infobars", "--disable-popup-blocking", "--window-position=0,0",
                          # nada de conversa com serviços do Google: mais rápido e menos ruído na rede da empresa
                          "--disable-background-networking", "--disable-component-update", "--disable-sync",
                          "--no-first-run", "--no-default-browser-check", "--disable-domain-reliability"])
        if self.mobile:
            base.update(APARELHO_ANDROID)
        else:
            base["viewport"] = {"width": 1366, "height": 850}
        base.update(extra)
        erros = []
        for canal in ("chrome", "msedge", "chromium", None):    # Chrome/Edge instalados passam melhor
            try:
                ctx = self._pw.chromium.launch_persistent_context(
                    str(self.perfil), **({"channel": canal} if canal else {}), **base)
                ctx.set_default_timeout(20000)
                ctx.add_init_script(JS_INIT)
                if self.mobile:
                    ctx.add_init_script(JS_INIT_ANDROID)
                self.aparelho = "android" if self.mobile else "desktop"
                log.debug(f"navegador aberto (canal: {canal or 'chromium do playwright'}, {self.aparelho})")
                return ctx
            except Exception as e:
                erros.append(f"{canal or 'chromium'}: {str(e).splitlines()[0] if str(e) else type(e).__name__}")
                if any(x in str(e) for x in ("ProcessSingleton", "already running", "profile is in use")):
                    raise NavegadorIndisponivel("feche outras execuções do script (perfil do navegador em uso)")
        if self.mobile:                                  # último recurso coerente: WebKit com UA de iPhone
            try:
                opcoes = {k: v for k, v in base.items() if k != "args"}
                opcoes.update(APARELHO_IPHONE)
                opcoes.update(extra)
                ctx = self._pw.webkit.launch_persistent_context(str(self.perfil) + "_wk", **opcoes)
                ctx.set_default_timeout(20000)
                ctx.add_init_script(JS_INIT)
                self.aparelho = "iphone"
                log.info("  (Chromium não abriu; usando WebKit com perfil de iPhone)")
                return ctx
            except Exception as e:
                erros.append(f"webkit: {str(e).splitlines()[0] if str(e) else type(e).__name__}")
        msg = " | ".join(erros)
        log.debug(f"nenhum canal de navegador abriu: {msg}")          # completo fica no log do dia
        if "Executable doesn't exist" in msg or "playwright install" in msg:
            raise NavegadorIndisponivel("falta instalar o navegador: python -m playwright install chromium"
                                        + (" webkit" if self.mobile else ""))
        if re.search(r"missing dependencies|Host system is missing", msg, re.I):
            raise NavegadorIndisponivel("faltam bibliotecas do sistema: python -m playwright install-deps chromium"
                                        " (no Colab/Linux, rode com sudo)")
        if "Missing X server" in msg or "$DISPLAY" in msg:
            raise NavegadorIndisponivel("esta máquina não tem tela; rode sem --ver (modo headless)")
        # a mensagem inteira está no log; aqui vai só o último canal tentado, que é o que importa
        raise NavegadorIndisponivel(f"não abriu o navegador: {erros[-1][:220] if erros else 'motivo desconhecido'}")

    def _pagina(self):
        return self.ctx.pages[0] if self.ctx.pages else self.ctx.new_page()

    def reinicia_pagina(self):
        try:
            self.page.close()
        except Exception:
            pass
        self.page = self.ctx.new_page()
        self._liga_captura()

    # ---- captura de XHR ---------------------------------------------------
    def _liga_captura(self):
        """Guarda as respostas JSON que a própria página busca: é lá que costuma estar preço + EAN."""
        def ao_responder(resp):
            try:
                tipo = (resp.headers or {}).get("content-type", "")
                if "json" in tipo.lower() and resp.request.resource_type in ("xhr", "fetch", "document", "other"):
                    self.respostas.append(resp)
                    if len(self.respostas) > 60:
                        del self.respostas[:-60]
            except Exception:
                pass
        try:
            self.page.on("response", ao_responder)
        except Exception:
            pass

    def limpa_captura(self):
        self.respostas = []

    def payloads(self) -> list:
        """[(url, objeto json)] das respostas capturadas; corpos já descartados são ignorados."""
        saida = []
        for resp in list(self.respostas):
            try:
                if int((resp.headers or {}).get("content-length") or 0) > 4_000_000:
                    continue
                saida.append((resp.url, resp.json()))
            except Exception:
                continue
        return saida

    def fechar(self):
        for f in (lambda: self.ctx and self.ctx.close(), lambda: self._pw.stop()):
            try:
                f()
            except Exception:
                pass
        self.ctx = None

    # ---- navegação básica -------------------------------------------------
    def _simular_humano(self):
        try:
            self.page.mouse.move(random.randint(200, 800), random.randint(200, 600))
            self.page.mouse.wheel(0, random.randint(150, 500))
            self.page.wait_for_timeout(random.randint(300, 700))
        except Exception:
            pass

    def _abrir(self, url: str) -> Optional[int]:
        """Abre a URL e devolve o HTTP status (None se não houve resposta). Levanta ErroRede / NavegadorMorreu."""
        self.navegacoes += 1
        if self.navegacoes % 60 == 0:       # página nova de vez em quando: evita inchar memória
            self.reinicia_pagina()
        codigo = None
        for tentativa in (1, 2):
            try:
                resp = self.page.goto(url, wait_until="domcontentloaded", timeout=45000)
                codigo = resp.status if resp else None
                break
            except self.TimeoutPW:
                log.debug(f"timeout abrindo {url}; seguindo com o que carregou")
                break
            except self.ErroPW as e:
                msg = str(e)
                if "ERR_ABORTED" in msg:
                    break                                   # navegação substituída por outra; segue
                if any(x in msg for x in ("Target closed", "has been closed", "crashed", "Connection closed")):
                    raise NavegadorMorreu(msg[:120])
                if "net::ERR_" in msg or "NS_ERROR" in msg or "chrome-error://" in msg:
                    if tentativa == 1:                      # a aba pode ter ficado presa numa página de erro
                        try:
                            self.page.goto("about:blank", timeout=10000)
                        except Exception:
                            pass
                        self.page.wait_for_timeout(1000)
                        continue
                    m = re.search(r"net::ERR_[A-Z_]+", msg)
                    raise ErroRede(m.group() if m else "página de erro do navegador")
                raise
        self._simular_humano()
        self._espera()
        self._fecha_popups()
        return codigo

    def _espera(self):
        try:
            self.page.wait_for_load_state("networkidle", timeout=6000)
        except Exception:
            pass
        self.page.wait_for_timeout(1000 + random.randint(200, 600))

    def _fecha_popups(self):
        for sel in BOTOES_POPUP:
            try:
                b = self.page.locator(sel).first
                if b.count() and b.is_visible():
                    b.click(timeout=1500)
                    self.page.wait_for_timeout(400)
            except Exception:
                pass
        try:
            self.page.keyboard.press("Escape")
        except Exception:
            pass

    def _bloqueado(self) -> bool:
        try:
            titulo = self.page.title() or ""
            texto = self.page.evaluate("() => (document.body && document.body.innerText || '').slice(0, 3000)") or ""
        except Exception:
            return False
        if RE_TITULO_BLOQ.search(titulo):
            return True
        if len(texto) < 1500 and RE_TEXTO_BLOQ.search(texto):
            return True
        try:
            # Desafios de página inteira. Não considera qualquer iframe contendo
            # "captcha": reCAPTCHA/hCaptcha invisível de login pode existir em
            # páginas normais, como na Nissei, sem bloquear catálogo ou preço.
            travas = self.page.locator(
                '#challenge-running, #challenge-form, #px-captcha, [id*="incapsula" i]'
            )
            if travas.first.count():
                return True

            frames = self.page.locator('iframe[src*="captcha" i]')
            for i in range(min(frames.count(), 5)):
                src = frames.nth(i).get_attribute("src") or ""
                if not re.search(r"recaptcha|gstatic|hcaptcha", src, re.I):
                    return True
            return False
        except Exception:
            return False

    def _html(self) -> str:
        try:
            return self.page.content()
        except Exception:
            self.page.wait_for_timeout(2000)
            return self.page.content()

    def _info(self) -> dict:
        try:
            return self.page.evaluate(JS_INFO_PAGINA) or {}
        except Exception:
            return {}

    def _debug(self, nome: str):
        try:
            PASTA_DEBUG.mkdir(parents=True, exist_ok=True)
            nome = re.sub(r"[^\w\-]", "_", nome)
            self.page.screenshot(path=str(PASTA_DEBUG / f"{nome}.png"), full_page=True)
            (PASTA_DEBUG / f"{nome}.html").write_text(self._html(), "utf-8")
        except Exception:
            pass

    # ---- busca ------------------------------------------------------------
    def _procura_input_visivel(self):
        for sel in SELETORES_BUSCA:
            try:
                loc = self.page.locator(sel)
                for i in range(min(loc.count(), 6)):
                    campo = loc.nth(i)
                    if campo.is_visible():
                        return campo
            except Exception:
                continue
        return None

    def _abre_busca_escondida(self) -> bool:
        """No mobile a caixa costuma ficar atrás de uma lupa; clica nela para revelar o campo."""
        for sel in BOTOES_ABRIR_BUSCA:
            try:
                loc = self.page.locator(sel)
                for i in range(min(loc.count(), 3)):
                    b = loc.nth(i)
                    if not b.is_visible():
                        continue
                    b.click(timeout=2500)
                    self.page.wait_for_timeout(700)
                    if self._procura_input_visivel() is not None:
                        return True
            except Exception:
                continue
        return False

    def _campo_busca(self):
        campo = self._procura_input_visivel()
        if campo is None and self._abre_busca_escondida():
            campo = self._procura_input_visivel()
        return campo

    def _pesquisar(self, cfg: dict, termo: str):
        """Faz a busca (URL direta ou caixa de busca). Devolve True ou um status de erro."""
        base = cfg["base"]
        dom = dominio(base)
        modelo = cfg.get("busca") or self.busca_aprendida.get(dom)
        self.url_digitada = None
        if modelo:
            if base not in self.aquecidos:                  # 1ª visita: home antes, para pegar cookies
                self._abrir(base)
                self.aquecidos.add(base)
                self.page.wait_for_timeout(1500)
            codigo = self._abrir(monta_url_busca(modelo, termo))
            if codigo is None or codigo < 400 or codigo == 404:
                return True
            log.debug(f"{dom}: busca por URL respondeu HTTP {codigo}; tentando pela caixa de busca")
            if not cfg.get("busca"):
                self.busca_aprendida.pop(dom, None)         # o modelo aprendido parou de funcionar
        codigo = self._abrir(base)
        self.aquecidos.add(base)
        if codigo and codigo >= 400:
            return f"SITE RESPONDEU HTTP {codigo} (tente --ver)"
        campo = self._campo_busca()
        if campo is None:
            return CAIXA_BUSCA
        try:
            campo.click(timeout=3000)
            self._simular_humano()
            campo.fill("")
            try:
                campo.press_sequentially(termo, delay=random.randint(60, 160))
            except AttributeError:                          # playwright antigo
                campo.type(termo, delay=random.randint(60, 160))
            self.page.wait_for_timeout(random.randint(300, 700))
            campo.press("Enter")
        except Exception as e:
            log.debug(f"{dom}: erro digitando na caixa de busca: {e}")
            return CAIXA_BUSCA
        self._espera()
        url = self.page.url
        if (termo.isdigit() and termo in url and RE_URL_BUSCA.search(url)
                and not self.busca_aprendida.get(dom)):
            self.url_digitada = url.replace(termo, "{termo}")
        return True

    def _aprende_busca(self, cfg: dict):
        """A busca digitada deu certo e a URL tem o termo: passa a usar a URL direta (bem mais rápido)."""
        if self.url_digitada and not cfg.get("busca"):
            dom = dominio(cfg["base"])
            self.busca_aprendida[dom] = self.url_digitada
            log.info(f"  ({dom}: aprendi a URL de busca -> {self.url_digitada})")
        self.url_digitada = None

    def _e_pdp(self, html: str, ean: str) -> bool:
        soup = BeautifulSoup(html or "", "html.parser")
        produtos, tipos = _produtos_jsonld(soup)
        if produtos:
            return not any(RE_TIPO_LISTA.search(t) for t in tipos)
        info = self._info()
        return bool(info.get("h1")) and 1 <= int(info.get("compra") or 0) <= 3 and not RE_URL_BUSCA.search(info.get("url") or "")

    def _preco_riscado(self, cands: list, preco: Optional[float]) -> Optional[float]:
        riscados = [p for p in (precos_no_texto(x.get("t", ""))[0] for x in cands if x.get("riscado")) if p]
        de = max(riscados) if riscados else None
        return de if de and preco and de > preco else None

    def _le_preco(self, html: str, ean: str) -> Optional[Resultado]:
        """Preço da página: estrutura (JSON-LD/meta) > JSON das APIs que a página chamou > DOM > texto."""
        res = extrai_preco_html(html, ean)
        try:
            cands = self.page.evaluate(JS_PRECO_DOM) or []
        except Exception:
            cands = []
        if res is not None and res.fonte != "texto":
            if res.preco_de is None:                     # completa o "de" riscado que o JSON-LD não traz
                res = replace(res, preco_de=self._preco_riscado(cands, res.preco))
            return res
        do_xhr, _ = preco_nas_capturas(self.payloads(), ean, res.preco if res else None, self._texto())
        if do_xhr is not None:
            if do_xhr.preco_de is None:
                do_xhr = replace(do_xhr, preco_de=self._preco_riscado(cands, do_xhr.preco))
            return do_xhr
        for c in cands:
            if c.get("riscado"):
                continue
            preco, _ = precos_no_texto(c.get("t", ""))
            if preco:
                return Resultado(OK, preco=preco, preco_de=self._preco_riscado(cands, preco), fonte="dom",
                                 nome_site=res.nome_site if res else None)
        return res

    def _texto(self) -> str:
        try:
            return self.page.evaluate("() => (document.body && document.body.innerText || '').slice(0, 20000)") or ""
        except Exception:
            return ""

    def _do_xhr(self, ean: str, url_pagina: str, base: Optional[str] = None) -> Optional[Resultado]:
        """Preço direto do JSON que a página buscou — serve mesmo sem abrir a página do produto."""
        res, url_api = preco_nas_capturas(self.payloads(), ean, None, self._texto())
        if res is None:
            return None
        url_produto = urljoin(base or url_pagina, res.url) if res.url else None
        if url_produto and dominio(url_produto) != dominio(base or url_pagina):
            url_produto = None                       # link para fora do site: não serve como página do produto
        return replace(res, url=url_produto or url_pagina,
                       obs="; ".join(x for x in [res.obs, f"preço lido da API do site ({dominio(url_api)})"] if x))

    def _le_pagina_produto(self, url: str, ean: str, nome) -> Optional[Resultado]:
        """Abre uma URL de produto conhecida e lê o preço, exigindo confirmação do EAN na página."""
        self.limpa_captura()
        self._abrir(url)
        if self._bloqueado():
            return Resultado(BLOQUEADO)
        html, atual = self._html(), self.page.url
        if not self._e_pdp(html, ean):
            return None
        res = self._le_preco(html, ean)
        if res is None:
            return None
        nivel, evid = confirma_ean_pagina(html, ean)
        if nivel == 0 and (res.evidencia or "") != "api da página:ean":
            sim = similaridade_nome(nome, res.nome_site or self._info().get("h1"))
            if sim < SIMILARIDADE_MIN:
                return None                     # a URL guardada não é (mais) deste produto: refaz a busca
        return classifica(res, nivel, evid, nome, res.nome_site or self._info().get("h1"), atual)

    def _url_conhecida(self, cfg: dict, ean: str) -> Optional[str]:
        if self.indice is None:
            return None
        return self.indice.url(cfg.get("base"), ean)

    def buscar(self, cfg: dict, ean: str, nome=None, debug: bool = False, cliente: str = "") -> Resultado:
        """Busca pelo EAN e, se preciso, pelo nome do SKU. Confirma o EAN na página e classifica a confiança."""
        # atalho: se já sabemos a página deste produto, vai direto nela (pula a busca, que é a parte bloqueada)
        conhecida = self._url_conhecida(cfg, ean)
        if conhecida:
            falha_de_rede = False
            try:
                res = self._le_pagina_produto(conhecida, ean, nome)
            except ErroRede:
                res, falha_de_rede = None, True          # rede ruim não invalida a URL guardada
            if res is not None:
                if res.achou or res.status == SEM_ESTOQUE:
                    return replace(res, obs="; ".join(x for x in [res.obs, "página já conhecida"] if x))
                if res.status == BLOQUEADO:
                    return res
            if not falha_de_rede:
                log.debug(f"URL guardada não serviu para {ean} em {dominio(cfg.get('base'))}; refazendo a busca")
                if self.indice is not None:
                    self.indice.esquece(cfg.get("base"), ean)

        termos = [("ean", ean)]
        tb = termo_busca(nome)
        if tb:
            termos.append(("nome", tb))
        melhor: Optional[Resultado] = None
        for tipo, termo in termos:
            self.limpa_captura()
            st = self._pesquisar(cfg, termo)
            if st is not True:
                return Resultado(st)
            if debug:
                self._debug(f"{cliente}_{ean}_{tipo}")
            if self._bloqueado():
                return Resultado(BLOQUEADO)
            html, url = self._html(), self.page.url
            pdp = self._e_pdp(html, ean)
            if not pdp:
                # a própria busca já trouxe o produto por JSON: não precisa abrir a página dele
                atalho = self._do_xhr(ean, url, cfg.get("base"))
                if atalho is not None and atalho.achou:
                    self._aprende_busca(cfg)
                    if self.indice is not None:
                        self.indice.aprende(cfg.get("base"), ean, atalho.url)
                    return replace(atalho, metodo=None)
                url = url_produto_jsonld(html, ean, cfg["base"])
                if not url:
                    try:
                        url = self.page.evaluate(JS_PRODUTO, ean)
                    except Exception:
                        url = None
                if not url:
                    continue
                self._abrir(url)
                if self._bloqueado():
                    return Resultado(BLOQUEADO)
                html, url = self._html(), self.page.url
                pdp = self._e_pdp(html, ean)
            res = self._le_preco(html, ean)
            if res is None:
                if debug:
                    self._debug(f"{cliente}_{ean}_{tipo}_produto")
                continue
            self._aprende_busca(cfg)
            nivel, evid = confirma_ean_pagina(html, ean)
            titulo = res.nome_site or self._info().get("h1") or None
            res = classifica(res, nivel, evid, nome, titulo, url)
            if not pdp and nivel == 0:                       # leu preço numa página que não parece ser de produto
                res = replace(res, status=OK_CONFERIR if res.status == OK else res.status, confianca="BAIXA",
                              obs="página aberta não parece ser de um produto")
            if res.confianca in ("ALTA", "MÉDIA"):
                if self.indice is not None and pdp:
                    self.indice.aprende(cfg.get("base"), ean, res.url)   # da próxima vez vai direto
                return res
            if melhor is None or ORDEM_CONFIANCA[res.confianca] > ORDEM_CONFIANCA[melhor.confianca]:
                melhor = res                                # BAIXA: ainda tenta pelo nome antes de aceitar

        if melhor is None and self.indice is not None:       # última carta: procurar o EAN no sitemap do site
            do_mapa = self.indice.url_no_sitemap(cfg.get("base"), ean)
            if do_mapa and do_mapa != conhecida:
                log.debug(f"tentando {do_mapa} (achado no sitemap)")
                try:
                    res = self._le_pagina_produto(do_mapa, ean, nome)
                except ErroRede:
                    res = None
                if res is not None and res.achou:
                    self.indice.aprende(cfg.get("base"), ean, res.url)
                    return replace(res, obs="; ".join(x for x in [res.obs, "página achada pelo sitemap"] if x))
        return melhor or Resultado(NAO_ENCONTRADO)


# ================================================================ motores ==
def candidatos(cfg: dict) -> list:
    """Métodos a tentar, na ordem, para um site."""
    motor = cfg.get("motor", "auto")
    lista = []
    if motor in ("vtex", "auto"):
        lista.append("vtex")
    if motor == "sfcc":
        lista.append("sfcc")
    return lista + ["nav", "nav_mobile", "nav_visivel"]


class Motores:
    """Executa um método num cliente, com navegador sob demanda e as demais bandeiras do grupo."""

    def __init__(self, visivel=False, debug=False, sem_navegador=False, busca_aprendida: Optional[dict] = None,
                 mobile=False, indice: Optional[IndiceUrls] = None, sufixo_perfil: str = ""):
        self.visivel_forcado, self.debug, self.sem_navegador = visivel, debug, sem_navegador
        self.sufixo_perfil = sufixo_perfil
        self.mobile_forcado = mobile
        self.busca_aprendida = busca_aprendida if busca_aprendida is not None else {}
        self.indice = indice
        self._nav: Optional[Navegador] = None
        # o erro é por perfil (headless / celular / visível): a janela visível falhar num servidor
        # sem tela NÃO pode derrubar o headless, que funcionava (era o bug das centenas de
        # "NAVEGADOR INDISPONÍVEL" depois do primeiro cliente que pedia janela visível)
        self._nav_erros: dict = {}
        self._nav_geral: Optional[str] = None        # navegador não abre de jeito nenhum (falta instalar)

    def navegador_fora(self) -> Optional[str]:
        """Motivo, se nenhum navegador consegue abrir nesta máquina. None = ainda vale tentar."""
        if self.sem_navegador:
            return "desligado por --sem-navegador"
        return self._nav_geral

    def nav(self, visivel: bool, mobile: bool = False) -> Navegador:
        if self.sem_navegador:
            raise NavegadorIndisponivel("desligado por --sem-navegador")
        if self._nav_geral:
            raise NavegadorIndisponivel(self._nav_geral)
        visivel = visivel or self.visivel_forcado
        mobile = mobile or self.mobile_forcado
        if visivel and not tem_tela():
            raise NavegadorIndisponivel("janela visível precisa de uma tela (sem DISPLAY nesta máquina)")
        chave = (visivel, mobile)
        if self._nav_erros.get(chave):
            raise NavegadorIndisponivel(self._nav_erros[chave])
        if self._nav is not None and (self._nav.visivel != visivel or self._nav.mobile != mobile):
            self._nav.fechar()                       # trocar de aparelho exige outro perfil/navegador
            self._nav = None
        if self._nav is None:
            try:
                self._nav = Navegador(visivel, self.busca_aprendida, mobile=mobile, indice=self.indice,
                                      sufixo_perfil=self.sufixo_perfil)
            except NavegadorIndisponivel as e:
                self._nav_erros[chave] = str(e)
                # falhou o perfil padrão da rodada (ou falta instalação): não abre nada, é geral.
                # qualquer outro motivo vale só para aquele perfil.
                padrao = (self.visivel_forcado and tem_tela(), self.mobile_forcado)
                if ERRO_INSTALACAO.search(str(e)) or chave in (padrao, (False, False)):
                    self._nav_geral = str(e)
                    log.warning(f"\n>>> Navegador indisponível: {e}\n"
                                f"    (clientes que só funcionam com navegador ficam sem preço nesta rodada)\n")
                else:
                    perfil = "janela visível" if visivel else "celular"
                    log.warning(f">>> Navegador em modo {perfil} indisponível: {e} — sigo com o modo normal")
                raise
        return self._nav

    def fechar(self):
        if self._nav:
            self._nav.fechar()
            self._nav = None

    def _roda(self, metodo: str, cfg: dict, cliente: str, ean: str, nome) -> Resultado:
        try:
            if metodo == "vtex":
                res = busca_vtex(cfg, ean, nome)
                if res.achou and self.indice is not None:
                    self.indice.aprende(cfg.get("base"), ean, res.url)
                return res
            if metodo == "sfcc":
                if not cfg.get("busca"):
                    return Resultado("SFCC SEM URL DE BUSCA")
                res = busca_sfcc(cfg, ean, nome, self.indice)
                if res.achou and self.indice is not None:
                    self.indice.aprende(cfg.get("base"), ean, res.url)
                return res
            if metodo in ("nav", "nav_mobile", "nav_visivel"):
                for tentativa in (1, 2):
                    try:
                        navegador = self.nav(metodo == "nav_visivel", mobile=metodo == "nav_mobile")
                        return navegador.buscar(cfg, ean, nome, self.debug, cliente)
                    except NavegadorIndisponivel:
                        return Resultado(NAV_INDISPONIVEL)
                    except ErroRede as e:
                        return Resultado(f"ERRO DE REDE: {e}")
                    except NavegadorMorreu as e:
                        log.debug(f"navegador caiu ({e}); reabrindo")
                        self.fechar()
                        if tentativa == 2:
                            raise
                    except Exception:
                        self.fechar()
                        if tentativa == 2:
                            raise
            return Resultado("MÉTODO DESCONHECIDO")
        except ErroRede as e:
            return Resultado(f"ERRO DE REDE: {e}")
        except Exception as e:
            log.debug(f"{cliente} {ean} {metodo}: exceção", exc_info=True)
            return Resultado(f"ERRO: {type(e).__name__}: {str(e)[:60]}")

    def _escada(self, metodo: str, cfg: dict, cliente: str, ean: str, nome, extras: int = 1) -> Resultado:
        """Tenta o método do dia e, se não resolver, escala para os próximos (API -> navegador -> celular)."""
        fila, res = [metodo] + [c for c in candidatos(cfg) if c != metodo], None
        usados = 0
        for cand in fila:
            if self.sem_navegador and cand.startswith("nav"):
                continue
            if cand == "nav_visivel" and (not self.visivel_forcado or not tem_tela()):
                continue                                   # janela visível é escolha do usuário, não automática
            if cand.startswith("nav") and self.navegador_fora():
                continue                                   # navegador nenhum abre: não gasta tentativa
            if cand == "nav_mobile" and self.mobile_forcado:
                continue                                   # com --mobile o "nav" já é celular
            r = self._roda(cand, cfg, cliente, ean, nome)
            r.metodo = cand
            r.bandeira = cfg.get("nome") or dominio(cfg.get("base"))
            if res is None or _nota_resultado(r) > _nota_resultado(res):
                res = r
            if r.achou or r.status == SEM_ESTOQUE:         # achou preço ou soube que está esgotado: encerrado
                break
            usados += 1
            if usados > extras:
                break
        if res is not None:
            return res
        # nenhum método pôde ser tentado (ex.: o cliente só funciona com navegador e ele está fora)
        return Resultado(NAV_INDISPONIVEL if self.navegador_fora() else "SEM MÉTODO APLICÁVEL", metodo=metodo)

    def executa(self, metodo: str, cliente: str, ean: str, nome=None, escalar: bool = False) -> Resultado:
        """Consulta a bandeira principal do cliente e, sem achar o EAN, tenta as outras bandeiras do grupo."""
        cfg = cfg_cliente(cliente)
        if not cfg or not cfg.get("base"):
            return Resultado(SEM_SITE)
        extras = 1 if escalar else 0

        def consulta(alvo: dict, metodo_alvo: str) -> Resultado:
            return (self._escada(metodo_alvo, alvo, cliente, ean, nome, extras) if escalar
                    else self._roda_simples(metodo_alvo, alvo, cliente, ean, nome))

        res = consulta(cfg, metodo)
        if res.achou or res.status == SEM_ESTOQUE:
            return res
        tentadas = []
        for banda in cfg.get("bandeiras") or []:
            m_banda = metodo if metodo in candidatos(banda) else candidatos(banda)[0]
            r2 = consulta(banda, m_banda)
            tentadas.append(banda.get("nome") or dominio(banda.get("base")))
            if r2.achou or r2.status == SEM_ESTOQUE:
                r2.obs = "; ".join(x for x in [r2.obs, f"veio da bandeira {banda.get('nome')}"] if x)
                return r2
        # nenhuma bandeira achou: fica valendo o resultado da principal (é o preço que interessa ao cliente)
        if tentadas:
            res.obs = "; ".join(x for x in [res.obs, f"também não achei em: {', '.join(tentadas)}"] if x)
        return res

    def _roda_simples(self, metodo: str, cfg: dict, cliente: str, ean: str, nome) -> Resultado:
        res = self._roda(metodo, cfg, cliente, ean, nome)
        res.metodo = metodo
        res.bandeira = cfg.get("nome") or dominio(cfg.get("base"))
        return res


class Metodos:
    """Método que passou em cada cliente (diagnóstico do dia) + URLs de busca aprendidas."""

    def __init__(self):
        self.clientes: dict = {}
        self.busca_aprendida: dict = {}
        for arq in (ARQ_METODOS, ARQ_METODOS_ANTIGO):
            if not arq.exists():
                continue
            try:
                d = json.loads(arq.read_text("utf-8"))
            except Exception:
                continue
            if isinstance(d, dict) and "clientes" in d:
                self.clientes = d.get("clientes") or {}
                self.busca_aprendida = d.get("busca_aprendida") or {}
            elif isinstance(d, dict):
                self.clientes = {k: v for k, v in d.items() if isinstance(v, dict)}
            break

    def metodo(self, cliente: str, padrao="nav") -> str:
        return (self.clientes.get(cliente) or {}).get("metodo") or padrao

    def salvar(self):
        try:
            PASTA.mkdir(parents=True, exist_ok=True)
            tmp = ARQ_METODOS.with_suffix(".tmp")
            tmp.write_text(json.dumps({"versao": 7, "clientes": dict(self.clientes),
                                       "busca_aprendida": dict(self.busca_aprendida)},
                                      ensure_ascii=False, indent=1), "utf-8")
            tmp.replace(ARQ_METODOS)
        except OSError as e:
            log.debug(f"não gravei {ARQ_METODOS}: {e}")


def diagnostico(alvos: pd.DataFrame, motores: Motores, metodos: Metodos, forcar: bool = False) -> Metodos:
    """Testa, com 1 EAN de cada cliente, qual método passa no site (uma vez por dia)."""
    hoje = str(date.today())
    pendentes = [c for c in alvos["CLIENTE"].unique()
                 if tem_site(c) and (forcar or (metodos.clientes.get(c) or {}).get("data") != hoje)]
    if not pendentes:
        return metodos
    log.info(f"\n=== Diagnóstico de {len(pendentes)} cliente(s): testando qual método passa em cada site ===")
    invalido = re.compile(r"NÃO É VTEX|CAIXA DE BUSCA|INDISPON|DESCONHECIDO|SEM URL")
    for cli in pendentes:
        cfg = cfg_cliente(cli)
        cands = [m for m in candidatos(cfg) if not (motores.sem_navegador and m.startswith("nav"))]
        if motores.mobile_forcado:
            cands = [m for m in cands if m != "nav_mobile"]      # --mobile já faz o "nav" ser mobile
        if not tem_tela():
            # sem tela a janela visível nem abre — tentar só estraga o navegador headless da rodada
            cands = [m for m in cands if m != "nav_visivel"]
        if not cands:
            log.info(f"  {cli[:25]:<25} -> (só funciona com navegador; pulado por --sem-navegador)")
            continue
        amostra = alvos[alvos["CLIENTE"] == cli].iloc[0]
        resultados = []
        for metodo in cands:
            if metodo == "nav_mobile" and any(m == "nav" and not r.temporario and not invalido.search(r.status)
                                              for m, r in resultados):
                continue                       # o desktop respondeu normalmente: só tenta o celular se bloquearem
            if metodo == "nav_visivel" and any(m.startswith("nav") and not r.temporario for m, r in resultados):
                break                          # o headless já respondeu (só não achou): não precisa do visível
            if metodo.startswith("nav") and motores.navegador_fora():
                break                          # já sabemos que navegador nenhum abre: não insiste
            res = motores.executa(metodo, cli, amostra["EAN"], amostra["SKU"])
            resultados.append((metodo, res))
            if res.achou:
                break
            if metodo == "vtex" and not invalido.search(res.status) and not res.temporario:
                break                          # a API é VTEX e respondeu (só não tem o EAN): fica no vtex
            time.sleep(3 + random.uniform(0, 1))
        com_preco = [m for m, r in resultados if r.achou]
        sem_bloqueio = [m for m, r in resultados if not r.temporario and not invalido.search(r.status)]
        if com_preco or sem_bloqueio:
            escolhido = (com_preco or sem_bloqueio)[0]
        elif any(r.temporario for _, r in resultados):
            escolhido = "bloqueado"
        elif motores.navegador_fora():
            # sem navegador não dá para concluir nada: não grava veredito, testa de novo quando ele voltar
            log.info(f"  {cli[:25]:<25} -> (sem método de API; precisa do navegador, que está fora)")
            continue
        else:
            escolhido = "sem_metodo"
        detalhe = " | ".join(f"{m}: {r.status}" for m, r in resultados)
        metodos.clientes[cli] = {"metodo": escolhido, "data": hoje, "teste": detalhe, "ean_teste": amostra["EAN"]}
        log.info(f"  {cli[:25]:<25} -> {escolhido:<12} ({detalhe})")
        metodos.salvar()
    motores.fechar()
    log.info("=== Fim do diagnóstico ===\n")
    return metodos


# ================================================================ agenda ==
@dataclass
class Item:
    cliente: str
    ean: str
    sku: Optional[str] = None
    tentativas: int = 0


class Agenda:
    """Alterna entre os sites respeitando um intervalo por domínio; recua e desiste quando bloqueiam."""

    def __init__(self, itens, metodos: Metodos, fator_ritmo: float = 1.0):
        self.filas: "OrderedDict[str, deque]" = OrderedDict()
        for it in itens:
            self.filas.setdefault(it.cliente, deque()).append(it)
        self.metodos = metodos
        self.fator_ritmo = max(1.0, fator_ritmo)
        self.site: dict = {}
        self.bloqueios: dict = {}
        self.sucessos: dict = {}
        self.abandonados: list = []
        self.duracoes: dict = {}

    @staticmethod
    def dominio_de(cli: str) -> str:
        return dominio((cfg_cliente(cli) or {}).get("base")) or cli

    def _estado(self, cli: str) -> dict:
        cfg = cfg_cliente(cli) or {}
        dom = self.dominio_de(cli)
        metodo = self.metodos.metodo(cli)
        ritmo = (cfg.get("ritmo") or RITMO_BASE.get(metodo, 5)) * self.fator_ritmo
        return self.site.setdefault(dom, {"intervalo": ritmo, "proximo": 0.0, "base_int": ritmo})

    def vazia(self) -> bool:
        return not any(self.filas.values())

    def descarta(self, cli: str) -> list:
        """Tira o cliente da rodada e devolve os EANs que sobraram na fila dele."""
        restantes = list(self.filas.get(cli) or [])
        if cli in self.filas:
            self.filas[cli].clear()
        return restantes

    def restantes(self) -> int:
        return sum(len(f) for f in self.filas.values())

    def proximo(self):
        while not self.vazia():
            agora = time.time()
            prontos = [c for c, f in self.filas.items() if f and self._estado(c)["proximo"] <= agora]
            if prontos:
                cli = prontos[0]
                self.filas.move_to_end(cli)
                return self.filas[cli].popleft()
            espera = min(self._estado(c)["proximo"] for c, f in self.filas.items() if f) - agora
            time.sleep(max(0.2, min(espera, 5)))
        return None

    def pega(self, ocupados: set) -> Optional[Item]:
        """Versão sem espera do proximo(), para o modo paralelo: um item de um site livre e no ritmo, ou None."""
        agora = time.time()
        for cli, fila in self.filas.items():
            if fila and self.dominio_de(cli) not in ocupados and self._estado(cli)["proximo"] <= agora:
                self.filas.move_to_end(cli)
                return fila.popleft()
        return None

    def espera(self, ocupados: set) -> float:
        """Segundos até algum site livre ficar pronto (0 se já há um; 5 se todos estão ocupados)."""
        livres = [self._estado(c)["proximo"] for c, f in self.filas.items()
                  if f and self.dominio_de(c) not in ocupados]
        return max(0.0, min(livres) - time.time()) if livres else 5.0

    def resultado(self, item: Item, res: Resultado, duracao: float = 0.0) -> str:
        """Devolve 'final', 'reagendado' ou 'desistiu'."""
        cli = item.cliente
        est = self._estado(cli)
        self.duracoes.setdefault(cli, []).append(duracao)
        if res.temporario:
            self.bloqueios[cli] = self.bloqueios.get(cli, 0) + 1
            limite = DESISTIR_APOS if not self.sucessos.get(cli) else DESISTIR_APOS * 3
            if self.bloqueios[cli] >= limite:
                self.abandonados = [(cli, it) for it in [item, *self.filas[cli]]]
                self.filas[cli].clear()
                return "desistiu"
            est["intervalo"] = min(est["intervalo"] * 2, RITMO_MAX)
            if item.tentativas + 1 < MAX_TENTATIVAS:
                item.tentativas += 1
                est["proximo"] = time.time() + ESFRIAR * item.tentativas * (2 if self.sucessos.get(cli) else 1)
                self.filas[cli].append(item)
                return "reagendado"
            est["proximo"] = time.time() + est["intervalo"]      # esgotou as tentativas: fica com o status
            return "final"
        self.bloqueios[cli] = 0
        self.sucessos[cli] = self.sucessos.get(cli, 0) + 1
        est["intervalo"] = max(est["base_int"], est["intervalo"] * 0.8)
        est["proximo"] = time.time() + est["intervalo"]
        return "final"

    def estimativa(self, paralelo: int = 1) -> Optional[float]:
        """Segundos restantes (estimativa grosseira): o maior entre o total de trabalho (dividido entre os
        trabalhadores) e o cliente mais lento."""
        total, mais_lento = 0.0, 0.0
        for cli, fila in self.filas.items():
            if not fila:
                continue
            d = self.duracoes.get(cli) or []
            media = (sum(d) / len(d)) if d else 8.0
            total += len(fila) * media
            mais_lento = max(mais_lento, len(fila) * (media + self._estado(cli)["intervalo"]))
        return max(total / max(1, paralelo), mais_lento) if (total or mais_lento) else None


# ============================================================== planilha ==
def acha_planilha(arg: Optional[str]) -> Path:
    if arg:
        p = Path(arg.strip('"'))
        if p.exists():
            return p
        parecidos = [x for x in Path(".").glob("*.xls*")
                     if not x.name.startswith("~$") and sem_acento(p.stem)[:6] in sem_acento(x.stem)]
        if len(parecidos) == 1:
            log.info(f"'{arg}' não existe, usando '{parecidos[0].name}'")
            return parecidos[0]
        sys.exit(f"\n>>> Arquivo '{arg}' não encontrado nesta pasta ({Path.cwd()}).\n"
                 f"    Planilhas aqui: {[x.name for x in Path('.').glob('*.xls*')] or 'nenhuma'}\n")
    def acha(*palavras):
        return sorted((x for x in Path(".").glob("*.xls*")
                       if any(w in sem_acento(x.name) for w in palavras) and not x.name.startswith("~$")
                       and not x.name.startswith("preco_site_")),
                      key=lambda x: x.stat().st_mtime, reverse=True)
    cands = acha("REGUA", "BUSCA") or acha("ACELERA")
    if not cands:
        sys.exit("\n>>> Nenhuma planilha de régua (Busca*/Régua*) nem 'Acelera' na pasta. "
                 "Informe: python preco_site_por_ean.py arquivo.xlsx\n")
    log.info(f"Usando planilha: {cands[0].name}")
    return cands[0]


def le_base(arquivo: Path):
    """Lê a aba Investimentos (ou a 1ª com CLIENTE e EAN). Devolve (alvos, ignorados)."""
    try:
        xls = pd.ExcelFile(arquivo)
    except PermissionError:
        sys.exit("\n>>> Não consegui abrir a planilha. Feche o arquivo no Excel e rode de novo.\n")
    except Exception as e:
        sys.exit(f"\n>>> Erro ao ler a planilha: {e}\n")

    def acha_colunas(colunas) -> dict:
        """{'CLIENTE': coluna real, 'EAN': ..., 'SKU': ...} aceitando variações ('CÓD EAN', 'GTIN', 'RAZÃO SOCIAL')."""
        cols = [sem_acento(c) for c in colunas]
        def pega(exatos, contem):
            for e in exatos:
                if e in cols:
                    return e
            return next((c for c in cols if any(x in c for x in contem)), None)
        return {"CLIENTE": pega(["CLIENTE"], ["CLIENTE", "RAZAO", "REDE", "CNPJ"]),
                "EAN": pega(["EAN"], ["EAN", "GTIN", "BARRAS"]),
                "SKU": pega(["SKU"], ["SKU", "DESCRI", "PRODUTO", "NOME", "ITEM"])}

    regua = le_regua(xls)
    if regua is not None:
        return regua

    aba = next((a for a in xls.sheet_names if sem_acento(a) == "INVESTIMENTOS"), None)
    if aba is None:
        for a in xls.sheet_names:
            try:
                achadas = acha_colunas(pd.read_excel(xls, a, nrows=0).columns)
            except Exception:
                continue
            if achadas["CLIENTE"] and achadas["EAN"]:
                aba = a
                break
    if aba is None:
        sys.exit(f"\n>>> Nenhuma aba com colunas CLIENTE e EAN. Abas: {xls.sheet_names}\n")

    df = pd.read_excel(xls, aba, dtype=object)
    df.columns = [sem_acento(c) for c in df.columns]
    achadas = acha_colunas(df.columns)
    faltando = [k for k in ("CLIENTE", "EAN") if not achadas[k]]
    if faltando:
        sys.exit(f"\n>>> A aba '{aba}' não tem a(s) coluna(s): {faltando}. Colunas: {list(df.columns)}\n")
    df = df.rename(columns={achadas["CLIENTE"]: "CLIENTE", achadas["EAN"]: "EAN"})
    if achadas["SKU"] and achadas["SKU"] != "SKU":
        df = df.rename(columns={achadas["SKU"]: "SKU"})
    if "SKU" not in df.columns:
        df["SKU"] = None
    df = df.loc[:, ~df.columns.duplicated()]
    df["SKU"] = df["SKU"].map(lambda v: "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v).strip())
    df["CLIENTE"] = df["CLIENTE"].map(lambda v: "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v).strip())
    df["EAN_OK"] = df["EAN"].map(normaliza_ean)

    tem_ean = df["EAN"].map(lambda v: v is not None and not (isinstance(v, float) and math.isnan(v)) and str(v).strip() != "")
    invalidos = df[df["EAN_OK"].isna() & tem_ean]
    ignorados = pd.DataFrame({"CLIENTE": invalidos["CLIENTE"], "EAN NA PLANILHA": invalidos["EAN"].astype(str),
                              "SKU": invalidos["SKU"], "MOTIVO": "EAN inválido (dígito verificador não confere)"})
    sem_cliente = df[df["EAN_OK"].notna() & (df["CLIENTE"] == "")]
    if len(sem_cliente):
        ignorados = pd.concat([ignorados, pd.DataFrame({"CLIENTE": "", "EAN NA PLANILHA": sem_cliente["EAN_OK"],
                                                        "SKU": sem_cliente["SKU"], "MOTIVO": "linha sem cliente"})])
    if len(ignorados):
        log.info(f"Aviso: {len(ignorados)} linha(s) ignorada(s) (EAN inválido ou sem cliente) -> aba 'Ignorados'")

    df = df[df["EAN_OK"].notna() & (df["CLIENTE"] != "")]
    alvos = (df[["CLIENTE", "EAN_OK", "SKU"]].rename(columns={"EAN_OK": "EAN"})
             .drop_duplicates(["CLIENTE", "EAN"]).reset_index(drop=True))
    return alvos, ignorados.reset_index(drop=True)


# colunas da régua que vão para o relatório: (nomes aceitos na planilha, nome no Excel)
COLS_REGUA = [
    (["TIPO"], "TIPO"),
    (["FAMILIA"], "FAMÍLIA"),
    (["SKU", "COD SKU", "CODIGO"], "CÓD. SKU"),
    (["APRESENTACAO", "DESCRICAO", "PRODUTO"], "APRESENTAÇÃO"),
]
COLUNA_PRECO_REGUA = 15           # coluna P da régua (0 = A): "Preço Sugerido Junho 26" — única referência de preço
FAIXAS = [                        # (desconto mínimo vs. sugerido, rótulo) — do maior para o menor
    (0.40, "MAIS DE 40% ABAIXO"), (0.25, "25% A 40% ABAIXO"), (0.10, "10% A 25% ABAIXO"),
    (0.01, "ATÉ 10% ABAIXO"), (-0.01, "NO SUGERIDO"),
]
EXTRAS_REGUA = [nome for _, nome in COLS_REGUA] + ["PREÇO SUGERIDO"]
RE_SUFIXO_APRES = re.compile(r"\s*-\s*(S|B|OFTA|[A-Z]{1,2})\s*$|\s+[SB]\s*$")


def clientes_regua() -> list:
    """Todos os clientes configurados que vendem online, sem repetir o mesmo site."""
    vistos, lista = set(), []
    for nome, cfg in SITES_TODOS.items():
        base = (cfg or {}).get("base")
        if not base or dominio(base) in vistos:
            continue
        vistos.add(dominio(base))
        lista.append(nome)
    return lista


def le_regua(xls: pd.ExcelFile):
    """Planilha de régua: tem EAN mas não tem CLIENTE; o cabeçalho pode não estar na 1ª linha.
    Devolve (alvos, ignorados) com cada EAN cruzado com todos os clientes, ou None se não for régua."""
    for aba in xls.sheet_names:
        try:
            topo = pd.read_excel(xls, aba, header=None, nrows=15, dtype=object)
        except Exception:
            continue
        linha_cab = next((i for i, lin in topo.iterrows()
                          if "EAN" in [sem_acento(v) for v in lin.tolist() if isinstance(v, str)]), None)
        if linha_cab is None:
            continue
        cab = [sem_acento(v) if isinstance(v, str) else "" for v in topo.iloc[linha_cab].tolist()]
        if any("CLIENTE" in c or "RAZAO" in c for c in cab):
            return None                                   # tem cliente: é o formato Acelera
        df = pd.read_excel(xls, aba, header=linha_cab, dtype=object)
        df.columns = [sem_acento(c) for c in df.columns]
        df = df.loc[:, ~df.columns.duplicated()]
        nome_col = next((c for c in ("APRESENTACAO", "DESCRICAO", "PRODUTO", "NOME") if c in df.columns), None)

        def texto(v):
            return "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v).strip()

        def num(v):
            try:
                f = float(v)
                return None if math.isnan(f) else round(f, 2)
            except (TypeError, ValueError):
                return None

        base = pd.DataFrame({"EAN_PLAN": df["EAN"],
                             "EAN": df["EAN"].map(normaliza_ean),
                             "SKU": (df[nome_col].map(texto).map(lambda t: RE_SUFIXO_APRES.sub("", t))
                                     if nome_col else "")})
        for aceitos, nome in COLS_REGUA:
            col = next((a for a in aceitos if a in df.columns), None)
            vals = df[col] if col else pd.Series([None] * len(df))
            base[nome] = vals.map(texto).values
        # preço de referência: SÓ a coluna P (confere pelo nome; se a coluna mudar de lugar, acha pelo nome)
        col_p = df.columns[COLUNA_PRECO_REGUA] if len(df.columns) > COLUNA_PRECO_REGUA else None
        if not (col_p and "SUGERIDO" in col_p and "MINIMO" not in col_p):
            col_p = next((c for c in df.columns if "SUGERIDO" in c and "MINIMO" not in c), col_p)
        base["PREÇO SUGERIDO"] = (df[col_p].map(num) if col_p else pd.Series([None] * len(df))).values
        log.info(f"Preço de referência da régua: coluna '{col_p}'")

        tem_ean = base["EAN_PLAN"].map(lambda v: texto(v) != "")
        sem_ean = int((~tem_ean & (base["SKU"] != "")).sum())
        if sem_ean:
            log.info(f"{sem_ean} linha(s) da régua sem EAN: descartadas")
        invalido = base[tem_ean & base["EAN"].isna()]
        ignorados = pd.DataFrame({"CLIENTE": "(todos)", "EAN NA PLANILHA": invalido["EAN_PLAN"].astype(str),
                                  "SKU": invalido["SKU"], "MOTIVO": "EAN inválido (dígito verificador não confere)"})

        eans = base[base["EAN"].notna()].drop_duplicates("EAN").drop(columns="EAN_PLAN")
        clientes = clientes_regua()
        alvos = (pd.DataFrame({"CLIENTE": clientes}).merge(eans, how="cross")
                 [["CLIENTE", "EAN", "SKU"] + EXTRAS_REGUA].reset_index(drop=True))
        log.info(f"Régua detectada (aba '{aba}'): {len(eans)} EAN(s) x {len(clientes)} cliente(s) "
                 f"= {len(alvos)} consultas")
        if len(ignorados):
            log.info(f"Aviso: {len(ignorados)} EAN(s) inválido(s) na régua -> aba 'Ignorados'")
        return alvos, ignorados
    return None


def _vazio(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v))


def posicao_regua(preco, sug) -> Optional[str]:
    """Faixa do preço do site em relação ao Preço Sugerido (desconto grande é possível: não é tratado como erro)."""
    if _vazio(preco) or not preco:
        return None
    if _vazio(sug) or not sug:
        return "SEM PREÇO NA RÉGUA"
    desconto = 1 - preco / sug
    for minimo, rotulo in FAIXAS:
        if desconto >= minimo:
            return rotulo
    return "ACIMA DO SUGERIDO"


def monta_regua_x_clientes(precos: pd.DataFrame) -> pd.DataFrame:
    """Uma linha por EAN, uma coluna por cliente (preço do site), + menor/mediana/maior e onde está o menor."""
    if precos.empty:
        return pd.DataFrame()
    chave = ["EAN", "TIPO", "FAMÍLIA", "CÓD. SKU", "APRESENTAÇÃO", "PREÇO SUGERIDO"]
    info = precos[chave].drop_duplicates("EAN")
    tabela = precos.pivot_table(index="EAN", columns="CLIENTE", values="PREÇO SITE", aggfunc="first", sort=False)
    clientes = [c for c in precos["CLIENTE"].drop_duplicates() if c in tabela.columns]
    tabela = tabela.reindex(columns=clientes)
    out = info.merge(tabela, left_on="EAN", right_index=True, how="left")
    vals = out[clientes] if clientes else pd.DataFrame(index=out.index)
    out["CLIENTES COM PREÇO"] = vals.notna().sum(axis=1)
    out["MENOR PREÇO"] = vals.min(axis=1)
    # linha a linha: EAN sem preço em nenhum cliente fica em branco (idxmin direto quebra no pandas 3)
    out["ONDE ESTÁ O MENOR"] = ([r.idxmin() if r.notna().any() else None for _, r in vals.iterrows()]
                                if clientes else None)
    out["MEDIANA SITES"] = vals.median(axis=1)
    out["MAIOR PREÇO"] = vals.max(axis=1)
    out["MEDIANA VS SUGERIDO"] = [round(m / s - 1, 4) if not _vazio(m) and not _vazio(s) and s else None
                                  for m, s in zip(out["MEDIANA SITES"], out["PREÇO SUGERIDO"])]
    fim = ["CLIENTES COM PREÇO", "MENOR PREÇO", "ONDE ESTÁ O MENOR", "MEDIANA SITES", "MAIOR PREÇO",
           "MEDIANA VS SUGERIDO"]
    out = out[chave + fim + clientes]
    return out.sort_values(["TIPO", "FAMÍLIA", "APRESENTAÇÃO"], na_position="last").reset_index(drop=True)


def filtra_clientes(alvos: pd.DataFrame, so: Optional[list]) -> pd.DataFrame:
    if not so:
        return alvos
    filtro = {sem_acento(c) for c in so}
    nomes = {c: sem_acento(c) for c in alvos["CLIENTE"].unique()}
    def bate(nome_norm):
        return any(f == nome_norm or re.search(rf"(?<![A-Z0-9]){re.escape(f)}(?![A-Z0-9])", nome_norm) for f in filtro)
    escolhidos = {c for c, n in nomes.items() if bate(n)}
    faltam = [f for f in filtro if not any(f == n or re.search(rf"(?<![A-Z0-9]){re.escape(f)}(?![A-Z0-9])", n)
                                          for n in nomes.values())]
    if faltam:
        log.info(f"Aviso: cliente(s) não encontrado(s) na planilha: {sorted(faltam)}")
    return alvos[alvos["CLIENTE"].isin(escolhidos)].reset_index(drop=True)


# ================================================================= cache ==
def cache_valido(registro: Optional[dict], agora: Optional[datetime] = None) -> bool:
    if not registro:
        return False
    agora = agora or datetime.now()
    carimbo = registro.get("coletado_em")
    if not carimbo:                                          # registro antigo (só data)
        return bool(registro.get("data") == str(date.today()) and registro.get("preco"))
    try:
        coletado = datetime.fromisoformat(carimbo)
    except (TypeError, ValueError):
        return False
    status = str(registro.get("status", ""))
    chave = "OK" if registro.get("preco") else next((k for k in CACHE_TTL_HORAS if k in status), None)
    horas = CACHE_TTL_HORAS.get(chave, 1)
    if horas <= 0:                                           # problema nosso, não do site: não vale guardar
        return False
    return agora - coletado < timedelta(hours=horas)


class Cache:
    """Resultados por CLIENTE|EAN com validade por status. Gravação atômica e no máximo a cada 5 s."""

    def __init__(self):
        self.itens: dict = {}
        self.ignorar: set = set()
        self._sujo, self._ultimo = False, 0.0
        for arq in (ARQ_CACHE, ARQ_CACHE_ANTIGO):
            if not arq.exists():
                continue
            try:
                d = json.loads(arq.read_text("utf-8"))
            except Exception:
                log.info(f"Aviso: cache {arq} corrompido, começando do zero")
                continue
            if isinstance(d, dict):
                self.itens = d.get("itens", d) if "itens" in d else {k: v for k, v in d.items() if isinstance(v, dict)}
            break

    @staticmethod
    def chave(cliente: str, ean: str) -> str:
        return f"{cliente}|{ean}"

    def registro(self, cliente: str, ean: str) -> Optional[dict]:
        return self.itens.get(self.chave(cliente, ean))

    def valido(self, cliente: str, ean: str) -> bool:
        k = self.chave(cliente, ean)
        return k not in self.ignorar and cache_valido(self.itens.get(k))

    def guarda(self, cliente: str, ean: str, res: Resultado, metodo: Optional[str] = None):
        k = self.chave(cliente, ean)
        self.itens[k] = {**res.dict(), "metodo": metodo or res.metodo,
                         "coletado_em": datetime.now().isoformat(timespec="seconds")}
        self.ignorar.discard(k)
        self._sujo = True
        self.salvar()

    def salvar(self, forcar: bool = False):
        if not self._sujo or (not forcar and time.time() - self._ultimo < 5):
            return
        try:
            PASTA.mkdir(parents=True, exist_ok=True)
            tmp = ARQ_CACHE.with_suffix(".tmp")
            tmp.write_text(json.dumps({"versao": 7, "itens": self.itens}, ensure_ascii=False, indent=1), "utf-8")
            tmp.replace(ARQ_CACHE)
            self._sujo, self._ultimo = False, time.time()
        except OSError as e:
            log.debug(f"não gravei o cache: {e}")


class IndiceUrls:
    """Guarda a URL do produto por (domínio, EAN).

    A página de busca é a mais protegida e a que mais erra. Sabendo a URL do produto, as rodadas
    seguintes vão direto na página dele — mais rápido, menos bloqueio e menos chance de pegar
    o produto errado. Aprende sozinho a cada acerto e completa pelo sitemap.xml do site.
    """

    def __init__(self):
        self.itens: dict = {}
        self._sujo = False
        self._sitemaps: dict = {}
        self._trava = threading.RLock()
        self._travas_sitemap: dict = {}      # uma por domínio: só um trabalhador baixa o sitemap de cada site
        try:
            if ARQ_INDICE.exists():
                d = json.loads(ARQ_INDICE.read_text("utf-8"))
                self.itens = d.get("itens", d) if isinstance(d, dict) else {}
        except Exception:
            log.debug(f"índice de URLs ilegível ({ARQ_INDICE}); começando vazio")

    def url(self, base: Optional[str], ean: str) -> Optional[str]:
        return ((self.itens.get(dominio(base)) or {}).get(ean) or {}).get("url")

    def aprende(self, base: Optional[str], ean: str, url: Optional[str]):
        dom = dominio(base)
        if not dom or not url or not url.startswith("http") or RE_URL_BUSCA.search(url):
            return                                   # só guarda página de produto, nunca de busca
        with self._trava:
            atual = self.url(base, ean)
            if atual == url:
                return
            self.itens.setdefault(dom, {})[ean] = {"url": url, "quando": str(date.today())}
            self._sujo = True

    def esquece(self, base: Optional[str], ean: str):
        with self._trava:
            if (self.itens.get(dominio(base)) or {}).pop(ean, None) is not None:
                self._sujo = True

    def salvar(self, forcar: bool = False):
        with self._trava:
            if not self._sujo and not forcar:
                return
            try:
                PASTA.mkdir(parents=True, exist_ok=True)
                tmp = ARQ_INDICE.with_suffix(".tmp")
                tmp.write_text(json.dumps({"versao": 7, "itens": self.itens}, ensure_ascii=False, indent=1), "utf-8")
                tmp.replace(ARQ_INDICE)
                self._sujo = False
            except OSError as e:
                log.debug(f"não gravei o índice de URLs: {e}")

    # ---- sitemap ----------------------------------------------------------
    def url_no_sitemap(self, base: Optional[str], ean: str) -> Optional[str]:
        """Procura no sitemap do site uma URL de produto que carregue o EAN (muitos sites põem no slug)."""
        urls = self._urls_sitemap(base)
        if not urls:
            return None
        canon = canonico_ean(ean)
        for u in urls:
            if ean in u or (len(canon) >= 8 and canon in u):
                return u
        return None

    def _urls_sitemap(self, base: Optional[str]) -> list:
        dom = dominio(base)
        if not dom or not base:
            return []
        with self._trava:
            trava_dom = self._travas_sitemap.setdefault(dom, threading.Lock())
        with trava_dom:
            return self._urls_sitemap_dom(base, dom)

    def _urls_sitemap_dom(self, base: str, dom: str) -> list:
        if dom in self._sitemaps:
            return self._sitemaps[dom]
        arq = PASTA_SITEMAPS / f"{re.sub(r'[^a-z0-9]+', '_', dom)}.txt"
        try:
            if arq.exists() and (date.today() - date.fromtimestamp(arq.stat().st_mtime)).days < DIAS_SITEMAP:
                urls = arq.read_text("utf-8").splitlines()
                self._sitemaps[dom] = urls
                return urls
        except OSError:
            pass
        urls = self._baixa_sitemap(base)
        self._sitemaps[dom] = urls
        if urls:
            try:
                PASTA_SITEMAPS.mkdir(parents=True, exist_ok=True)
                arq.write_text("\n".join(urls), "utf-8")
            except OSError:
                pass
        return urls

    @staticmethod
    def _baixa_sitemap(base: str) -> list:
        """Baixa o sitemap (e os filhos) do site, com teto de tamanho e de tempo."""
        raiz = [f"{base}/sitemap.xml", f"{base}/sitemap_index.xml", f"{base}/sitemap/sitemap.xml"]
        vistos, fila, urls, inicio = set(), list(raiz), [], time.time()
        while fila and len(urls) < MAX_URLS_SITEMAP and time.time() - inicio < 120:
            alvo = fila.pop(0)
            if alvo in vistos:
                continue
            vistos.add(alvo)
            try:
                r = get(alvo)
            except ErroRede:
                continue
            if r.status_code != 200 or "xml" not in (r.headers.get("content-type") or "").lower():
                continue
            achadas = re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", r.text, re.I)
            for u in achadas:
                if u.endswith(".xml") or "sitemap" in u.lower().rsplit("/", 1)[-1]:
                    if len(vistos) + len(fila) < 60:
                        fila.append(u)
                elif dominio(u) == dominio(base):
                    urls.append(u)
        if urls:
            log.info(f"  ({dominio(base)}: {len(urls)} URLs lidas do sitemap)")
        return urls[:MAX_URLS_SITEMAP]


class Historico:
    """historico_precos.csv: uma linha por preço coletado. Alimenta 'PREÇO ANTERIOR' e 'VARIAÇÃO' no Excel."""
    COLUNAS = ["data", "hora", "cliente", "ean", "preco", "preco_de", "status", "metodo", "url"]

    def __init__(self):
        self._por_chave: Optional[dict] = None

    def registra(self, cliente: str, ean: str, res: Resultado):
        if not res.achou:
            return
        try:
            PASTA.mkdir(parents=True, exist_ok=True)
            novo = not ARQ_HISTORICO.exists()
            agora = datetime.now()
            with ARQ_HISTORICO.open("a", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                if novo:
                    w.writerow(self.COLUNAS)
                w.writerow([agora.strftime("%Y-%m-%d"), agora.strftime("%H:%M:%S"), cliente, ean,
                            f"{res.preco:.2f}".replace(".", ","),
                            f"{res.preco_de:.2f}".replace(".", ",") if res.preco_de else "",
                            res.status, res.metodo or "", res.url or ""])
            if self._por_chave is not None:
                self._por_chave.setdefault((cliente, ean), []).append(
                    (agora.strftime("%Y-%m-%d"), agora.strftime("%H:%M:%S"), res.preco))
        except OSError as e:
            log.debug(f"não gravei o histórico: {e}")

    def _carrega(self):
        self._por_chave = {}
        if not ARQ_HISTORICO.exists():
            return
        try:
            with ARQ_HISTORICO.open(encoding="utf-8-sig", newline="") as f:
                for linha in csv.DictReader(f, delimiter=";"):
                    try:
                        preco = float(str(linha.get("preco", "")).replace(",", "."))
                    except ValueError:
                        continue
                    self._por_chave.setdefault((linha.get("cliente"), linha.get("ean")), []).append(
                        (linha.get("data", ""), linha.get("hora", ""), preco))
        except OSError:
            pass

    @staticmethod
    def _mais_recente(regs: list):
        """Escolhe pela data/hora da coleta — não pelo maior preço (era o efeito de max() na tupla)."""
        d, _h, p = max(regs, key=lambda r: (r[0], r[1]))
        return p, d

    def anterior(self, cliente: str, ean: str, antes_de: str):
        """Último preço registrado em data anterior a `antes_de` (AAAA-MM-DD)."""
        if self._por_chave is None:
            self._carrega()
        regs = [r for r in self._por_chave.get((cliente, ean), []) if r[0] and r[0] < antes_de]
        return self._mais_recente(regs) if regs else (None, None)

    def ultimo(self, cliente: str, ean: str):
        """Preço mais recente já coletado, de qualquer data — usado quando a coleta de hoje falhou."""
        if self._por_chave is None:
            self._carrega()
        regs = [r for r in self._por_chave.get((cliente, ean), []) if r[0]]
        return self._mais_recente(regs) if regs else (None, None)


# ============================================================= relatório ==
def _mediana(vals: list) -> float:
    v = sorted(vals)
    n = len(v)
    return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2


def monta_relatorio(alvos: pd.DataFrame, ignorados: pd.DataFrame, cache: Cache, historico: Historico,
                    metodos: Metodos) -> dict:
    """Monta os DataFrames das abas Preços, Conferir, Resumo e Ignorados."""
    linhas, por_ean = [], {}
    for cli, ean, sku in alvos[["CLIENTE", "EAN", "SKU"]].itertuples(index=False):
        reg = cache.registro(cli, ean) or {}
        preco = reg.get("preco")
        status = reg.get("status") or (NAO_CONSULTADO if tem_site(cli) else SEM_SITE)
        coletado = str(reg.get("coletado_em") or "")
        anterior, data_ant = (historico.anterior(cli, ean, coletado[:10] or str(date.today()))
                              if preco else (None, None))
        variacao = round((preco - anterior) / anterior, 4) if preco and anterior else None
        # sem preço hoje: mostra o último já coletado, em coluna própria e com data (nunca na coluna do dia)
        ultimo, data_ultimo = (None, None) if preco else historico.ultimo(cli, ean)
        linhas.append({
            "CLIENTE": cli, "EAN": ean, "PREÇO SITE": preco, "SKU": sku, "STATUS": status,
            "CONFIANÇA": reg.get("confianca"), "PRODUTO NO SITE": reg.get("nome_site"),
            "PREÇO \"DE\"": reg.get("preco_de"), "PREÇO ANTERIOR": anterior, "VARIAÇÃO": variacao,
            "DATA ANTERIOR": data_ant, "ÚLTIMO PREÇO CONHECIDO": ultimo, "DATA DO ÚLTIMO": data_ultimo,
            "BANDEIRA": reg.get("bandeira"), "EVIDÊNCIA": reg.get("evidencia"), "OBS": reg.get("obs"),
            "MÉTODO": reg.get("metodo"), "FONTE": reg.get("fonte"),
            "COLETADO EM": coletado.replace("T", " "), "URL": reg.get("url"),
        })
        if preco:
            por_ean.setdefault(ean, []).append((cli, preco))
    precos = pd.DataFrame(linhas)
    tem_regua = all(c in alvos.columns for c in EXTRAS_REGUA) and len(precos)
    if tem_regua:
        extra = alvos[["CLIENTE", "EAN"] + EXTRAS_REGUA].drop_duplicates(["CLIENTE", "EAN"])
        precos = precos.merge(extra, on=["CLIENTE", "EAN"], how="left")
        precos["DIF. VS SUGERIDO"] = [
            round(p / s - 1, 4) if not _vazio(p) and not _vazio(s) and p and s else None
            for p, s in zip(precos["PREÇO SITE"], precos["PREÇO SUGERIDO"])]
        precos["POSIÇÃO NA RÉGUA"] = [posicao_regua(p, s) for p, s in zip(precos["PREÇO SITE"], precos["PREÇO SUGERIDO"])]
        primeiro = ["CLIENTE", "EAN", "TIPO", "FAMÍLIA", "APRESENTAÇÃO", "PREÇO SITE", "PREÇO SUGERIDO",
                    "DIF. VS SUGERIDO", "POSIÇÃO NA RÉGUA", "STATUS", "CONFIANÇA", "PRODUTO NO SITE", "PREÇO \"DE\""]
        precos = precos[primeiro + [c for c in precos.columns if c not in primeiro and c not in ("SKU", "CÓD. SKU")]
                        + ["CÓD. SKU"]]

    conferir = []
    for lin in precos.to_dict("records"):
        p, ean = lin["PREÇO SITE"], lin["EAN"]
        p = None if p is None or (isinstance(p, float) and math.isnan(p)) else p
        if not p:
            continue
        outros = [v for c, v in por_ean.get(ean, []) if c != lin["CLIENTE"]]
        motivo, med = None, None
        if len(outros) >= 2:
            med = _mediana(outros)
            razao = p / med if med else 1
            if razao < 0.25 and not tem_regua:
                motivo = "muito abaixo dos outros clientes: provável leitura errada"
            elif razao > 1.7 and not (tem_regua and not _vazio(lin.get("PREÇO SUGERIDO"))
                                      and p <= lin["PREÇO SUGERIDO"] * 1.15):   # no sugerido não é problema
                motivo = "bem acima dos outros clientes (sem ação? outra embalagem?)"
        if lin["CONFIANÇA"] == "BAIXA":
            motivo = (motivo + " | " if motivo else "") + "EAN não confirmado na página (conferir SKU)"
        var = lin["VARIAÇÃO"]
        if var is not None and not (isinstance(var, float) and math.isnan(var)) and abs(var) >= 0.30:
            motivo = (motivo + " | " if motivo else "") + f"variou {var:+.0%} vs. coleta anterior"
        if motivo:
            conferir.append({"CLIENTE": lin["CLIENTE"], "EAN": ean, "SKU": lin.get("APRESENTAÇÃO") or lin.get("SKU"), "PREÇO LIDO": p,
                             "MEDIANA OUTROS": round(med, 2) if med else None, "PRODUTO NO SITE": lin["PRODUTO NO SITE"],
                             "MOTIVO": motivo, "URL": lin["URL"]})
    conferir = pd.DataFrame(conferir, columns=["CLIENTE", "EAN", "SKU", "PREÇO LIDO", "MEDIANA OUTROS",
                                               "PRODUTO NO SITE", "MOTIVO", "URL"])

    resumo = []
    for cli, grupo in precos.groupby("CLIENTE", sort=False):
        st = grupo["STATUS"].fillna("")
        com_preco = int(grupo["PREÇO SITE"].notna().sum())
        cfg = cfg_cliente(cli) or {}
        tem_loja = bool(cfg.get("base"))
        bandas = [cfg.get("nome") or dominio(cfg.get("base"))] + [b.get("nome") or dominio(b.get("base"))
                                                                 for b in cfg.get("bandeiras") or []]
        resumo.append({
            "CLIENTE": cli, "EANS": len(grupo), "COM PREÇO": com_preco,
            # cobertura só faz sentido para quem tem loja online; sem loja fica em branco, não zero
            "COBERTURA": (round(com_preco / len(grupo), 4) if len(grupo) else 0) if tem_loja else None,
            "SEM ESTOQUE": int((st == SEM_ESTOQUE).sum()),
            "NÃO ENCONTRADO": int(st.isin([NAO_ENCONTRADO, SEM_PRECO]).sum()),
            "BLOQUEADO / ERRO": int(st.map(lambda s: bool(TEMPORARIO.search(s)) or s in (SITE_BLOQUEOU, SEM_METODO)).sum()),
            "CONFERIR": int((grupo["CONFIANÇA"] == "BAIXA").sum()),
            "NÃO CONSULTADO": int(st.isin([NAO_CONSULTADO, SEM_SITE, NAV_INDISPONIVEL]).sum()),
            **({"ABAIXO DO SUGERIDO": int(grupo["POSIÇÃO NA RÉGUA"].fillna("").str.contains("ABAIXO").sum()),
                "NO SUGERIDO": int((grupo["POSIÇÃO NA RÉGUA"] == "NO SUGERIDO").sum()),
                "ACIMA DO SUGERIDO": int((grupo["POSIÇÃO NA RÉGUA"] == "ACIMA DO SUGERIDO").sum()),
                "DESCONTO MEDIANO": (lambda d: round(-float(d.median()), 4) if len(d) else None)(
                    grupo["DIF. VS SUGERIDO"].dropna())}
               if tem_regua else {}),
            "MÉTODO": metodos.metodo(cli, "") if tem_loja else "",
            "BANDEIRAS": " > ".join(b for b in bandas if b) if tem_loja else "",
            "SITE": cfg.get("base") or "(não vende online)",
        })
    resumo = pd.DataFrame(resumo)
    if len(resumo):
        com_loja = resumo[resumo["COBERTURA"].notna()]
        geral = {"CLIENTE": "TOTAL (quem vende online)", "EANS": int(com_loja["EANS"].sum()),
                 "COM PREÇO": int(com_loja["COM PREÇO"].sum()),
                 "COBERTURA": round(com_loja["COM PREÇO"].sum() / com_loja["EANS"].sum(), 4) if len(com_loja) and com_loja["EANS"].sum() else None,
                 "SEM ESTOQUE": int(com_loja["SEM ESTOQUE"].sum()), "NÃO ENCONTRADO": int(com_loja["NÃO ENCONTRADO"].sum()),
                 "BLOQUEADO / ERRO": int(com_loja["BLOQUEADO / ERRO"].sum()), "CONFERIR": int(com_loja["CONFERIR"].sum()),
                 "NÃO CONSULTADO": int(com_loja["NÃO CONSULTADO"].sum()), "MÉTODO": "", "BANDEIRAS": "", "SITE": ""}
        if tem_regua:
            for c in ("ABAIXO DO SUGERIDO", "NO SUGERIDO", "ACIMA DO SUGERIDO"):
                geral[c] = int(com_loja[c].sum())
            d = precos["DIF. VS SUGERIDO"].dropna()
            geral["DESCONTO MEDIANO"] = round(-float(d.median()), 4) if len(d) else None
        resumo = pd.concat([resumo, pd.DataFrame([geral])], ignore_index=True)
    abas = {"Preços": precos}
    if tem_regua:
        abas["Régua x Clientes"] = monta_regua_x_clientes(precos)
    abas.update({"Conferir": conferir, "Resumo": resumo, "Ignorados": ignorados})
    return abas


LARGURAS = {
    "Preços": {"CLIENTE": 26, "EAN": 16, "PREÇO SITE": 12, "SKU": 34, "STATUS": 20, "CONFIANÇA": 11,
               "PRODUTO NO SITE": 44, "PREÇO \"DE\"": 11, "PREÇO ANTERIOR": 12, "VARIAÇÃO": 10, "DATA ANTERIOR": 13,
               "ÚLTIMO PREÇO CONHECIDO": 15, "DATA DO ÚLTIMO": 13, "BANDEIRA": 20,
               "EVIDÊNCIA": 26, "OBS": 34, "MÉTODO": 11, "FONTE": 9, "COLETADO EM": 19, "URL": 60,
               "TIPO": 10, "FAMÍLIA": 18, "CÓD. SKU": 10, "APRESENTAÇÃO": 36, "PREÇO SUGERIDO": 12,
               "DIF. VS SUGERIDO": 11, "POSIÇÃO NA RÉGUA": 22},
    "Régua x Clientes": {"EAN": 16, "TIPO": 10, "FAMÍLIA": 18, "CÓD. SKU": 10, "APRESENTAÇÃO": 36,
                         "PREÇO SUGERIDO": 12, "CLIENTES COM PREÇO": 10, "MENOR PREÇO": 12, "ONDE ESTÁ O MENOR": 20,
                         "MEDIANA SITES": 12, "MAIOR PREÇO": 12, "MEDIANA VS SUGERIDO": 12},
    "Conferir": {"CLIENTE": 26, "EAN": 16, "SKU": 34, "PREÇO LIDO": 12, "MEDIANA OUTROS": 15,
                 "PRODUTO NO SITE": 44, "MOTIVO": 60, "URL": 60},
    "Resumo": {"CLIENTE": 26, "EANS": 8, "COM PREÇO": 11, "COBERTURA": 11, "SEM ESTOQUE": 12, "NÃO ENCONTRADO": 15,
               "BLOQUEADO / ERRO": 16, "CONFERIR": 10, "NÃO CONSULTADO": 15, "MÉTODO": 11,
               "ABAIXO DO SUGERIDO": 15, "NO SUGERIDO": 12, "ACIMA DO SUGERIDO": 15, "DESCONTO MEDIANO": 13,
               "BANDEIRAS": 40, "SITE": 44},
    "Ignorados": {"CLIENTE": 26, "EAN NA PLANILHA": 22, "SKU": 34, "MOTIVO": 50},
}
FORMATOS = {"R$": {"PREÇO SITE", "PREÇO \"DE\"", "PREÇO ANTERIOR", "PREÇO LIDO", "MEDIANA OUTROS",
                   "ÚLTIMO PREÇO CONHECIDO", "PREÇO SUGERIDO",
                   "MENOR PREÇO", "MEDIANA SITES", "MAIOR PREÇO"},
            "%": {"VARIAÇÃO", "COBERTURA", "DIF. VS SUGERIDO", "MEDIANA VS SUGERIDO", "DESCONTO MEDIANO"}}


def _formata_aba(ws, df: pd.DataFrame, aba: str):
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = "G2" if aba == "Régua x Clientes" else "A2"
    ws.auto_filter.ref = ws.dimensions
    verde, vermelho = PatternFill("solid", fgColor="E2F0D9"), PatternFill("solid", fgColor="FCE4E4")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for i, col in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = LARGURAS.get(aba, {}).get(col, 14)
        if len(df) == 0:
            continue
        celulas = [c[0] for c in ws.iter_rows(min_row=2, max_row=len(df) + 1, min_col=i, max_col=i)]
        if col in FORMATOS["R$"] or (aba == "Régua x Clientes" and i > 12):   # colunas de cliente = preço
            for c in celulas:
                c.number_format = '"R$" #,##0.00'
        elif col in FORMATOS["%"]:
            for c in celulas:
                c.number_format = "0.0%"
                # variação: verde caiu de preço, vermelho subiu, sem cor quando ficou igual ou está vazio.
                # (nunca ler c.fill para "manter como está": openpyxl devolve um StyleProxy que não pode ser reatribuído)
                if col == "VARIAÇÃO" and isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    if c.value < 0:
                        c.fill = verde
                    elif c.value > 0:
                        c.fill = vermelho
        elif col == "URL":
            for c in celulas:
                if isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink, c.style = c.value, "Hyperlink"
        elif col == "POSIÇÃO NA RÉGUA":
            for c in celulas:
                if c.value == "NO SUGERIDO":
                    c.fill = verde
                elif c.value in ("25% A 40% ABAIXO", "MAIS DE 40% ABAIXO"):
                    c.fill = vermelho
                elif c.value == "ACIMA DO SUGERIDO":
                    c.fill = amarelo
        elif col == "CONFIANÇA":
            for c in celulas:
                if c.value == "BAIXA":
                    c.fill = vermelho


ARQ_PARCIAL = Path("preco_site_PARCIAL.xlsx")


def salva_excel(abas: dict, parcial: bool = False) -> Optional[Path]:
    nome = str(ARQ_PARCIAL) if parcial else f"preco_site_{datetime.now():%Y%m%d_%H%M}.xlsx"
    for n in range(1 if parcial else 20):
        destino = Path(nome if n == 0 else nome.replace(".xlsx", f"_{n}.xlsx"))
        if destino.exists() and not parcial:      # rodou de novo no mesmo minuto: não sobrescreve
            continue
        try:
            with pd.ExcelWriter(destino, engine="openpyxl") as xw:
                for aba, df in abas.items():
                    if aba != "Preços" and len(df) == 0:
                        continue
                    df.to_excel(xw, index=False, sheet_name=aba)
                    # a formatação é enfeite: se der errado, o Excel sai sem cor, mas sai
                    try:
                        _formata_aba(xw.sheets[aba], df, aba)
                    except Exception:
                        log.debug(f"não formatei a aba '{aba}'", exc_info=True)
                        log.warning(f"  (aviso: a aba '{aba}' saiu sem formatação)")
            return destino
        except PermissionError:
            continue
        except Exception:
            log.debug(f"falhei ao gravar {destino}", exc_info=True)
            try:
                destino.unlink()                  # não deixa arquivo pela metade na pasta
            except OSError:
                pass
            raise
    log.warning("\n>>> Não consegui salvar o Excel (feche arquivos abertos). Os preços continuam no cache.")
    return None


def gera_relatorio(alvos, ignorados, cache, historico, metodos, parcial: bool = False) -> Optional[Path]:
    """Monta as abas, salva o Excel e imprime a cobertura por cliente.

    parcial=True: regrava preco_site_PARCIAL.xlsx no meio da rodada, sem imprimir o resumo."""
    abas = monta_relatorio(alvos, ignorados, cache, historico, metodos)
    destino = salva_excel(abas, parcial=parcial)
    precos = abas["Preços"]
    if parcial:
        if destino:
            log.info(f"  (Excel parcial atualizado: {destino} | com preço até agora: "
                     f"{int(precos['PREÇO SITE'].notna().sum())} de {len(precos)})")
        return destino
    if destino:
        try:
            ARQ_PARCIAL.unlink()                  # o final substitui o parcial
        except OSError:
            pass
        log.info(f"\nSalvo em {destino}  |  com preço: {int(precos['PREÇO SITE'].notna().sum())} de {len(precos)}"
                 f"  |  para conferir: {len(abas['Conferir'])}  |  ignorados: {len(ignorados)}")
    imprime_resumo(abas["Resumo"])
    return destino


def imprime_resumo(resumo: pd.DataFrame):
    if resumo is None or len(resumo) == 0:
        return
    log.info("\nCobertura por cliente:")
    for _, lin in resumo.iterrows():
        extras = [f"{col.lower()}: {int(lin[col])}"
                  for col in ("SEM ESTOQUE", "NÃO ENCONTRADO", "BLOQUEADO / ERRO", "CONFERIR", "NÃO CONSULTADO")
                  if lin[col]]
        cob = lin["COBERTURA"]
        pct = f"({float(cob):>5.0%})" if cob is not None and not pd.isna(cob) else "(não vende online)"
        log.info(f"  {str(lin['CLIENTE'])[:25]:<25} {int(lin['COM PREÇO']):>4}/{int(lin['EANS']):<4} {pct}"
                 + ("  " + " | ".join(extras) if extras else ""))


# ================================================================== main ==
def _fmt_tempo(seg: Optional[float]) -> str:
    if seg is None:
        return "?"
    seg = int(seg)
    return f"{seg // 3600}h{(seg % 3600) // 60:02d}min" if seg >= 3600 else f"{seg // 60}min" if seg >= 60 else f"{seg}s"


def consulta_item(motores: "Motores", metodos: "Metodos", item: Item, escalar: bool = False):
    """Consulta um EAN num cliente (com o plano B pelo navegador). Devolve (método do dia, Resultado)."""
    cli = item.cliente
    metodo = metodos.metodo(cli)
    res = motores.executa(metodo, cli, item.ean, item.sku, escalar=escalar)
    if (not escalar and metodo in ("vtex", "sfcc") and not motores.navegador_fora()
            and (cfg_cliente(cli) or {}).get("motor") != "vtex"
            and (res.status in (NAO_ENCONTRADO, SEM_PRECO, NAO_E_VTEX) or res.temporario)):
        r2 = motores.executa("nav", cli, item.ean, item.sku)   # plano B: navegador (busca por EAN e nome)
        # o plano B só substitui a resposta da API se souber mais: "navegador indisponível"
        # não pode apagar um "NÃO ENCONTRADO" que o site já respondeu
        if r2.status != NAV_INDISPONIVEL and (r2.achou or not r2.temporario):
            res = r2
    return metodo, res


class Rodada:
    """O que acontece depois de cada consulta: agenda, cache, histórico, log e Excel parcial.

    Roda sempre na thread principal — os trabalhadores em paralelo só consultam e devolvem o resultado,
    então cache e histórico nunca são gravados por duas threads ao mesmo tempo.
    """

    def __init__(self, itens: list, metodos, cache, historico, indice, fator_ritmo: float, paralelo: int,
                 parcial=None):
        self.agenda = Agenda(itens, metodos, fator_ritmo)
        self.cache, self.historico, self.indice = cache, historico, indice
        self.feitos, self.total, self.paralelo = 0, len(itens), paralelo
        self.sem_preco: list = []
        self.parcial = parcial                     # função que regrava o Excel parcial (ou None)
        self._ultimo_parcial = time.time()

    def trata(self, item: Item, metodo: str, res: Resultado, duracao: float, nav_fora: Optional[str]):
        cli, agenda, cache = item.cliente, self.agenda, self.cache
        # navegador fora do ar e este cliente só funciona com ele: não faz sentido repetir
        # o mesmo erro EAN por EAN — encerra o cliente de uma vez
        if res.status == NAV_INDISPONIVEL and metodo.startswith("nav") and nav_fora:
            restantes = agenda.descarta(cli)
            cache.guarda(cli, item.ean, res)
            self.feitos += 1
            for it in restantes:
                cache.guarda(cli, it.ean, res)
                self.feitos += 1
            log.info(f"  !! {cli}: precisa do navegador, que está fora -> {1 + len(restantes)} EAN(s) "
                     f"sem preço nesta rodada ({nav_fora})")
            return
        decisao = agenda.resultado(item, res, duracao)
        if decisao == "desistiu":
            n = len(agenda.abandonados)
            log.info(f"  !! {cli}: {agenda.bloqueios[cli]} bloqueios seguidos -> site está barrando o robô. "
                     f"Pulando os {n} EAN(s) restante(s) dele.")
            for c, it in agenda.abandonados:
                cache.guarda(c, it.ean, Resultado(SITE_BLOQUEOU, metodo=res.metodo))
                self.sem_preco.append(it)
                self.feitos += 1
            agenda.abandonados = []
            return
        if decisao == "reagendado":
            espera = ESFRIAR * item.tentativas * (2 if agenda.sucessos.get(cli) else 1)
            log.info(f"    | {cli[:22]:<22} | {item.ean} | {res.metodo or metodo:<11} | {res.status}"
                     f"  -> tenta de novo em ~{max(1, espera // 60)} min")
            return
        self.feitos += 1
        cache.guarda(cli, item.ean, res)
        self.historico.registra(cli, item.ean, res)
        if self.indice is not None:
            self.indice.salvar()
        # vale reprocessar depois? só o que pode mudar de resultado numa nova tentativa
        sem_saida = (res.status in (SEM_ESTOQUE, SEM_SITE)
                     or (res.status == NAV_INDISPONIVEL and nav_fora))
        if not res.achou and not sem_saida:
            self.sem_preco.append(item)
        extra = f"  [{res.confianca}]" if res.confianca and res.confianca != "ALTA" else ""
        eta = (f"  (~{_fmt_tempo(agenda.estimativa(self.paralelo))} restantes)"
               if self.feitos % 10 == 0 and agenda.restantes() else "")
        banda = f" | {res.bandeira}" if res.bandeira and res.bandeira != (cfg_cliente(cli) or {}).get("nome") else ""
        log.info(f"{self.feitos:>4}/{self.total} | {cli[:22]:<22} | {item.ean} | {res.metodo or metodo:<11} | "
                 f"{res.status} | {res.preco if res.preco is not None else ''}{banda}{extra}{eta}")
        self._talvez_parcial()

    def _talvez_parcial(self):
        if self.parcial is None or time.time() - self._ultimo_parcial < PARCIAL_MIN * 60:
            return
        self._ultimo_parcial = time.time()
        try:
            self.cache.salvar(forcar=True)
            self.parcial()
        except Exception:
            log.debug("não gravei o Excel parcial", exc_info=True)


def roda_fila(itens: list, motores: "Motores", metodos: "Metodos", cache: "Cache", historico: "Historico",
              indice: Optional[IndiceUrls] = None, escalar: bool = False, fator_ritmo: float = 1.0,
              paralelo: int = 1, fabrica_motores=None, parcial=None) -> list:
    """Percorre a fila alternando entre sites. Devolve os itens que terminaram sem preço.

    paralelo > 1: vários sites ao mesmo tempo, cada um no seu ritmo e com no máximo UMA consulta
    em andamento por site — o site não vê diferença; a rodada termina muito antes.
    """
    rodada = Rodada(itens, metodos, cache, historico, indice, fator_ritmo, paralelo, parcial)
    agenda = rodada.agenda
    paralelo = max(1, min(paralelo, len(agenda.filas)))
    log.info(f"Consultando {rodada.total} EAN(s) em {len(agenda.filas)} cliente(s), "
             + (f"{paralelo} site(s) ao mesmo tempo...\n" if paralelo > 1 else "alternando entre sites...\n"))
    if paralelo <= 1 or fabrica_motores is None:
        while True:
            item = agenda.proximo()
            if item is None:
                break
            t0 = time.time()
            metodo, res = consulta_item(motores, metodos, item, escalar)
            rodada.trata(item, metodo, res, time.time() - t0, motores.navegador_fora())
        return rodada.sem_preco

    pedidos, respostas = queue.Queue(), queue.Queue()

    def trabalhador(n: int):
        # cada trabalhador tem seus motores: navegador (Playwright) e sessão HTTP próprios, na sua thread
        m = fabrica_motores(n)
        try:
            while True:
                item = pedidos.get()
                if item is None:
                    break
                t0 = time.time()
                try:
                    metodo, res = consulta_item(m, metodos, item, escalar)
                except Exception as e:
                    log.debug(f"trabalhador {n}: {item.cliente} {item.ean}", exc_info=True)
                    metodo, res = metodos.metodo(item.cliente), Resultado(f"ERRO: {type(e).__name__}: {str(e)[:60]}")
                respostas.put((item, metodo, res, time.time() - t0, m.navegador_fora()))
        finally:
            m.fechar()

    threads = [threading.Thread(target=trabalhador, args=(n,), daemon=True, name=f"trab{n}")
               for n in range(paralelo)]
    for t in threads:
        t.start()
    ocupados: set = set()              # domínios com consulta em andamento
    try:
        while True:
            while len(ocupados) < paralelo:
                item = agenda.pega(ocupados)
                if item is None:
                    break
                ocupados.add(agenda.dominio_de(item.cliente))
                pedidos.put(item)
            if not ocupados and agenda.vazia():
                break
            try:
                resposta = respostas.get(timeout=max(0.2, min(agenda.espera(ocupados), 2.0)))
            except queue.Empty:
                continue
            ocupados.discard(agenda.dominio_de(resposta[0].cliente))
            rodada.trata(*resposta)
    finally:
        while True:                    # interrompido: não começa mais nada, só deixa terminar o que está em curso
            try:
                pedidos.get_nowait()
            except queue.Empty:
                break
        for _ in threads:
            pedidos.put(None)
        for t in threads:
            t.join(timeout=60)
        while True:                    # respostas que chegaram depois da interrupção: não jogar fora
            try:
                rodada.trata(*respostas.get_nowait())
            except queue.Empty:
                break
    return rodada.sem_preco


def main(args) -> int:
    configura_log()
    carrega_sites_extra()
    log.info(f"Preço site por EAN v9  |  {datetime.now():%d/%m/%Y %H:%M}  |  pasta: {Path.cwd()}")
    alvos, ignorados = le_base(acha_planilha(args.arquivo))
    alvos = filtra_clientes(alvos, args.so)
    if args.limite:
        alvos = alvos.groupby("CLIENTE", sort=False).head(args.limite).reset_index(drop=True)
    if alvos.empty:
        log.info("\n>>> Nada para consultar.\n")
        return 1

    sem_site = sorted({c for c in alvos["CLIENTE"] if not tem_site(c)})
    if sem_site:
        log.info(f"Aviso: sem site configurado (ficam sem preço): {sem_site}")
    log.info(f"{len(alvos)} EAN(s) em {alvos['CLIENTE'].nunique()} cliente(s)")

    cache, historico, metodos, indice = Cache(), Historico(), Metodos(), IndiceUrls()
    if args.tudo:
        cache.ignorar = {Cache.chave(c, e) for c, e in zip(alvos["CLIENTE"], alvos["EAN"])}
        log.info("Modo: RODAR TUDO — ignora o cache e consulta os EANs desta seleção de novo")
    else:
        log.info("Modo: CONTINUAR — consulta só o que falta ou já venceu no cache (use --tudo para refazer tudo)")
    if args.relatorio:
        gera_relatorio(alvos, ignorados, cache, historico, metodos)
        return 0

    if args.ver and not tem_tela():
        log.warning("Aviso: --ver pede uma janela e esta máquina não tem tela (Colab/servidor). "
                    "Continuo em modo headless.")
        args.ver = False
    if not args.sem_navegador:
        falta = checa_navegador()
        if falta:
            log.warning(f"\n>>> Navegador indisponível nesta máquina: {falta}\n"
                        f"    Sites que só leem com navegador ficam sem preço. Para resolver:\n"
                        f"      pip install playwright && python -m playwright install chromium\n"
                        f"      (Linux/Colab, se reclamar de biblioteca: python -m playwright install-deps chromium)\n"
                        f"    Para rodar assim mesmo, sem avisos, use --sem-navegador.\n")

    motores = Motores(args.ver, args.debug, args.sem_navegador, metodos.busca_aprendida,
                      mobile=args.mobile, indice=indice)
    if args.mobile:
        log.info("Modo celular ligado: navegador emulando um Pixel 8 (Android/Chrome)")
    inicio, codigo_saida = time.time(), 0
    try:
        pend = [Item(c, e, s) for c, e, s in alvos[["CLIENTE", "EAN", "SKU"]].itertuples(index=False)
                if tem_site(c) and not cache.valido(c, e)]
        if not pend:
            log.info("Tudo já coletado (cache ainda válido). Use --tudo para consultar de novo.")
        else:
            chaves = {(it.cliente, it.ean) for it in pend}
            base_pend = alvos[[(c, e) in chaves for c, e in zip(alvos["CLIENTE"], alvos["EAN"])]]
            diagnostico(base_pend, motores, metodos, args.rediagnosticar)

            pulados = {}
            for it in pend:
                m = metodos.metodo(it.cliente)
                if m in ("bloqueado", "sem_metodo"):
                    pulados[it.cliente] = m
            if pulados:
                barrados = sorted(c for c, m in pulados.items() if m == "bloqueado")
                sem_metodo = sorted(c for c, m in pulados.items() if m == "sem_metodo")
                if barrados:
                    log.info(f"!! Bloquearam o robô em todos os métodos no teste (pulados nesta rodada): {barrados}")
                if sem_metodo:
                    log.info(f"!! Nenhum método leu o site no teste (pulados nesta rodada; veja {ARQ_METODOS}): {sem_metodo}")
                for it in pend:
                    if it.cliente in pulados:
                        teste = (metodos.clientes.get(it.cliente) or {}).get("teste")
                        cache.guarda(it.cliente, it.ean,
                                     Resultado(SITE_BLOQUEOU if pulados[it.cliente] == "bloqueado" else SEM_METODO, obs=teste))
                pend = [it for it in pend if it.cliente not in pulados]

            paralelo = max(1, args.paralelo)

            def fabrica_motores(n: int) -> Motores:
                # trabalhador 0 usa o perfil de sempre (cookies já aquecidos); os demais, perfis próprios
                return Motores(args.ver, args.debug, args.sem_navegador, metodos.busca_aprendida,
                               mobile=args.mobile, indice=indice, sufixo_perfil=f"_w{n}" if n else "")

            def parcial():
                gera_relatorio(alvos, ignorados, cache, historico, metodos, parcial=True)

            extra = dict(paralelo=paralelo, fabrica_motores=fabrica_motores, parcial=parcial)
            sobraram = roda_fila(pend, motores, metodos, cache, historico, indice, **extra)
            for n_passada in range(2, max(1, args.passadas) + 1):
                if not sobraram:
                    break
                log.info(f"\n=== {n_passada}ª passada: {len(sobraram)} EAN(s) sem preço, tentando outros "
                         f"caminhos e mais devagar ===")
                for it in sobraram:
                    it.tentativas = 0
                sobraram = roda_fila(sobraram, motores, metodos, cache, historico, indice,
                                     escalar=True, fator_ritmo=1.0 + 0.8 * (n_passada - 1), **extra)
    except KeyboardInterrupt:
        log.info("\nInterrompido. Salvando o que já foi coletado...")
        codigo_saida = 130
    except Exception:
        log.exception("\n>>> Erro inesperado. Salvando o que já foi coletado...")
        codigo_saida = 1
    finally:
        motores.fechar()
        metodos.salvar()
        cache.salvar(forcar=True)
        indice.salvar(forcar=True)
        try:
            gera_relatorio(alvos, ignorados, cache, historico, metodos)
        except Exception:
            log.exception(">>> Erro ao gerar o Excel (os preços continuam no cache; tente --relatorio)")
            codigo_saida = codigo_saida or 1
        log.info(f"\nTempo total: {_fmt_tempo(time.time() - inicio)}  |  log: {PASTA_LOGS}")
    return codigo_saida


def listar_clientes():
    configura_log()
    carrega_sites_extra()
    metodos = Metodos()
    log.info(f"{'CLIENTE':<28} {'MOTOR':<10} {'MÉTODO HOJE':<12} SITE")
    for nome, cfg in SITES_NORM.items():
        m = metodos.clientes.get(nome) or {}
        hoje = m.get("metodo", "") if m.get("data") == str(date.today()) else ""
        alts = "".join(f"  (+ {a.get('nome') or dominio(a['base'])})" for a in cfg.get("bandeiras") or [])
        log.info(f"{nome:<28} {cfg['motor']:<10} {hoje:<12} {cfg.get('base') or '(sem site)'}{alts}")
    if metodos.busca_aprendida:
        log.info("\nURLs de busca aprendidas:")
        for d, u in metodos.busca_aprendida.items():
            log.info(f"  {d:<35} {u}")


def testar(args) -> int:
    configura_log()
    carrega_sites_extra()
    if len(args.testar) < 2:
        sys.exit("Uso: --testar CLIENTE EAN")
    *cli, ean_bruto = args.testar
    cliente = " ".join(cli)
    ean = normaliza_ean(ean_bruto)
    if not ean:
        sys.exit(f"EAN inválido: {ean_bruto}")
    cfg = cfg_cliente(cliente)
    if not cfg or not cfg.get("base"):
        sys.exit(f"Cliente '{cliente}' sem site configurado. Clientes: {sorted(SITES_NORM)}")
    metodos = Metodos()
    m = Motores(args.ver, args.debug, args.sem_navegador, metodos.busca_aprendida, mobile=args.mobile)
    try:
        for metodo in candidatos(cfg):
            if args.sem_navegador and metodo.startswith("nav"):
                break
            if metodo == "nav_mobile" and args.mobile:
                continue
            r = m.executa(metodo, cliente, ean, args.nome)
            log.info(f"{metodo:<12} {r.dict()}")
            if r.achou:
                break
    finally:
        m.fechar()
        metodos.salvar()
    return 0


def mapear_api(args) -> int:
    """Abre o site, busca o EAN e lista os JSONs que a página pediu — mostra onde o site guarda o preço."""
    configura_log()
    carrega_sites_extra()
    if len(args.mapear_api) < 2:
        sys.exit("Uso: --mapear-api CLIENTE EAN")
    *cli, ean_bruto = args.mapear_api
    cliente = " ".join(cli)
    ean = normaliza_ean(ean_bruto)
    if not ean:
        sys.exit(f"EAN inválido: {ean_bruto}")
    cfg = cfg_cliente(cliente)
    if not cfg or not cfg.get("base"):
        sys.exit(f"Cliente '{cliente}' sem site configurado. Clientes: {sorted(SITES_NORM)}")

    metodos = Metodos()
    nav = Navegador(args.ver, metodos.busca_aprendida, mobile=args.mobile)
    achados = []
    try:
        log.info(f"Mapeando {dominio(cfg['base'])} ({nav.aparelho}) com o EAN {ean}...")
        res = nav.buscar(cfg, ean, args.nome, args.debug, cliente)
        log.info(f"Resultado da busca: {res.dict()}\n")
        for url, dados in nav.payloads():
            tem_ean = ean in json.dumps(dados, ensure_ascii=False, default=str)
            preco = preco_em_json(dados, ean)
            achados.append({"url": url, "tem_ean": tem_ean,
                            "preco_encontrado": preco.preco if preco else None,
                            "nome": preco.nome_site if preco else None})
        uteis = [a for a in achados if a["preco_encontrado"] is not None]
        quase = [a for a in achados if a["tem_ean"] and a["preco_encontrado"] is None]
        log.info(f"{len(achados)} resposta(s) JSON capturada(s); {len(uteis)} com EAN + preço:\n")
        for a in uteis:
            log.info(f"  [PREÇO {a['preco_encontrado']}] {a['url'][:160]}")
        for a in quase:
            log.info(f"  [tem o EAN, sem preço legível] {a['url'][:160]}")
        if not uteis and not quase:
            log.info("  (nenhuma resposta JSON trouxe o EAN — esse site provavelmente renderiza o preço no HTML)")
        try:
            PASTA.mkdir(parents=True, exist_ok=True)
            destino = PASTA / f"apis_{re.sub(r'[^a-z0-9]+', '_', dominio(cfg['base']))}.json"
            destino.write_text(json.dumps({"cliente": cliente, "ean": ean, "aparelho": nav.aparelho,
                                           "quando": datetime.now().isoformat(timespec="seconds"),
                                           "respostas": achados}, ensure_ascii=False, indent=1), "utf-8")
            log.info(f"\nDetalhes salvos em {destino}")
        except OSError as e:
            log.warning(f"não gravei o mapa de APIs: {e}")
    finally:
        nav.fechar()
        metodos.salvar()
    return 0


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Preço site por EAN — consulta o preço de cada EAN no e-commerce de cada cliente")
    p.add_argument("arquivo", nargs="?", help="régua de preços (EANs sem cliente) ou Acelera (CLIENTE|EAN|SKU); sem nome, acha sozinho na pasta")
    p.add_argument("--so", nargs="+", metavar="CLIENTE", help="só estes clientes")
    p.add_argument("--testar", nargs="+", metavar="CLIENTE_EAN", help="ex.: --testar PAGUE MENOS 7896004707037")
    p.add_argument("--mapear-api", nargs="+", metavar="CLIENTE_EAN",
                   help="mostra os JSONs que o site chama (onde está o preço): --mapear-api NISSEI 7896004707037")
    p.add_argument("--nome", help="nome do SKU para o --testar / --mapear-api")
    p.add_argument("--ver", action="store_true", help="navegador visível")
    p.add_argument("--mobile", action="store_true",
                   help="emula celular (Android/Chrome): HTML mais simples e menos bloqueio")
    p.add_argument("--novo", "--tudo", dest="tudo", action="store_true",
                   help="roda tudo: ignora o cache e consulta todos os EANs da seleção de novo")
    p.add_argument("--passadas", type=int, default=2, metavar="N",
                   help="quantas passadas fazer; da 2ª em diante só reprocessa o que ficou sem preço (padrão: 2)")
    p.add_argument("--limite", type=int, metavar="N", help="só N EANs por cliente (teste rápido)")
    p.add_argument("--paralelo", type=int, default=PARALELO_PADRAO, metavar="N",
                   help=f"quantos sites consultar ao mesmo tempo, cada um no seu ritmo (padrão: {PARALELO_PADRAO}; "
                        "1 = um de cada vez, como antes)")
    p.add_argument("--relatorio", action="store_true", help="só regera o Excel a partir do cache")
    p.add_argument("--sem-navegador", action="store_true", help="não abre navegador (só VTEX/SFCC)")
    p.add_argument("--debug", action="store_true", help=f"salva prints e HTML em {PASTA_DEBUG}/")
    p.add_argument("--rediagnosticar", action="store_true", help="refaz o teste de método por cliente")
    p.add_argument("--listar", action="store_true", help="mostra os clientes configurados e sai")
    return p


if __name__ == "__main__":
    a = parser().parse_args()
    if a.listar:
        listar_clientes()
    elif a.mapear_api:
        sys.exit(mapear_api(a))
    elif a.testar:
        sys.exit(testar(a))
    else:
        sys.exit(main(a))
